"""Belegtext und Seitenbilder für das KI-Auslesen (synthetische Dokumente)."""
import base64
import zipfile
from datetime import date
from email.message import EmailMessage
from io import BytesIO

import pytest

from app.services import pdf_kontoauszug
from app.services.beleg_text import (
    DOCX,
    EML,
    PDF,
    XLSX,
    beleg_text,
    mime_bestimmen,
    vision_bilder,
)
from app.services.pdf_kontoauszug import ocr_verfuegbar

ocr_noetig = pytest.mark.skipif(not ocr_verfuegbar(), reason="tesseract nicht installiert")

ZEILEN = [
    "Bürobedarf Schulz GmbH · Hauptstraße 1 · 12345 Berlin",
    "Auftragsbestätigung Nr. AB-4711",
    "Pos 1   Kopierpapier A4, 50 Kartons          1.234,56",
    "Summe netto 1.234,56 EUR",
]


def _text_pdf(zeilen: list[str], seiten: int = 1) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    puffer = BytesIO()
    c = canvas.Canvas(puffer, pagesize=A4)
    for nr in range(seiten):
        y = 800
        for zeile in [*zeilen, f"Seite {nr + 1}"]:
            c.drawString(40, y, zeile)
            y -= 16
        c.showPage()
    c.save()
    return puffer.getvalue()


def _seitenbild(zeilen: list[str], dpi: int = 200):
    import pypdfium2 as pdfium

    dokument = pdfium.PdfDocument(_text_pdf(zeilen))
    try:
        return dokument[0].render(scale=dpi / 72).to_pil()
    finally:
        dokument.close()


def _bild_pdf(zeilen: list[str]) -> bytes:
    """Seite rendern und nur als Bild wieder einbetten (Scan ohne Textebene)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    breite, hoehe = A4
    puffer = BytesIO()
    c = canvas.Canvas(puffer, pagesize=A4)
    c.drawImage(ImageReader(_seitenbild(zeilen)), 0, 0, width=breite, height=hoehe)
    c.showPage()
    c.save()
    return puffer.getvalue()


def _png(bild) -> bytes:
    puffer = BytesIO()
    bild.save(puffer, "PNG")
    return puffer.getvalue()


def _docx() -> bytes:
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    dokument = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document {w}><w:body>
  <w:p><w:r><w:t>Bestellung Nr. B-2026-17</w:t></w:r></w:p>
  <w:p><w:r><w:t xml:space="preserve">An: </w:t></w:r><w:r><w:t>Metallbau Kraft KG</w:t></w:r></w:p>
  <w:tbl>
    <w:tr><w:tc><w:p><w:r><w:t>Pos</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Artikel</w:t></w:r></w:p></w:tc>
          <w:tc><w:p><w:r><w:t>Betrag</w:t></w:r></w:p></w:tc></w:tr>
    <w:tr><w:tc><w:p><w:r><w:t>1</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Stahlträger</w:t></w:r></w:p></w:tc>
          <w:tc><w:p><w:r><w:t>2.500,00</w:t></w:r></w:p></w:tc></w:tr>
  </w:tbl>
  <w:p><w:r><w:t>Summe</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>2.500,00 EUR</w:t></w:r></w:p>
</w:body></w:document>"""
    kopf = f"""<?xml version="1.0" encoding="UTF-8"?>
<w:hdr {w}><w:p><w:r><w:t>Nordlicht Handel GmbH · Hafenstraße 5</w:t></w:r></w:p></w:hdr>"""
    puffer = BytesIO()
    with zipfile.ZipFile(puffer, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", dokument)
        z.writestr("word/header1.xml", kopf)
        z.writestr("word/header2.xml", kopf)   # gleiche Kopfzeile in zweitem Abschnitt
    return puffer.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    mappe = Workbook()
    blatt = mappe.active
    blatt.title = "Bestellung"
    blatt.append(["Lieferant", "Holz Meier GmbH"])
    blatt.append(["Datum", date(2026, 11, 3)])
    blatt.append(["Artikel", "Menge", "Preis"])
    blatt.append(["Eichenbretter", 40, 12.5])
    blatt.append(["Summe", None, 500.0])
    mappe.create_sheet("Leer")
    puffer = BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def _eml(anhang: bytes | None = None, html: bool = False) -> bytes:
    nachricht = EmailMessage()
    nachricht["Subject"] = "Bestellung 4711 – Büromöbel"
    nachricht["From"] = "einkauf@nordlicht.example"
    nachricht["To"] = "vertrieb@moebel-hansen.example"
    if html:
        nachricht.set_content("<html><body><p>Sehr geehrte Damen und Herren,</p>"
                              "<table><tr><td>Schreibtisch</td><td>899,00</td></tr></table>"
                              "<script>alert(1)</script></body></html>", subtype="html")
    else:
        nachricht.set_content("Hiermit bestellen wir 3 Schreibtische zu je 899,00 EUR netto.")
    if anhang is not None:
        nachricht.add_attachment(anhang, maintype="application", subtype="pdf",
                                 filename="bestellung.pdf")
    return bytes(nachricht)


def _bild_aus_data_url(url: str):
    from PIL import Image

    assert url.startswith("data:image/png;base64,")
    return Image.open(BytesIO(base64.b64decode(url.split(",", 1)[1])))


# ------------------------------------------------------------------ PDF

def test_pdf_mit_textebene():
    erg = beleg_text(_text_pdf(ZEILEN), PDF)
    assert erg.quelle == "TEXT"
    assert erg.seiten == 1
    assert "Auftragsbestätigung Nr. AB-4711" in erg.text
    assert "1.234,56" in erg.text
    assert "        " not in erg.text      # Layout-Leerraum verdichtet
    assert erg.warnungen == []


def test_pdf_nur_erste_seiten():
    erg = beleg_text(_text_pdf(ZEILEN, seiten=7), PDF, max_seiten=5)
    assert erg.seiten == 7
    assert "Seite 5" in erg.text and "Seite 6" not in erg.text
    assert any("5 von 7 Seiten" in w for w in erg.warnungen)


def test_pdf_ohne_textebene_ohne_ocr(monkeypatch):
    monkeypatch.setattr(pdf_kontoauszug, "ocr_verfuegbar", lambda: False)
    erg = beleg_text(_bild_pdf(ZEILEN), PDF)
    assert erg.quelle == "LEER"
    assert erg.text == ""
    assert any("tesseract" in w for w in erg.warnungen)


@ocr_noetig
def test_pdf_ohne_textebene_per_ocr():
    erg = beleg_text(_bild_pdf(ZEILEN), PDF)
    assert erg.quelle == "OCR", erg.warnungen
    assert "AB-4711" in erg.text
    assert "1.234,56" in erg.text
    assert any("Texterkennung" in w for w in erg.warnungen)


def test_defekte_dateien_ohne_ausnahme(monkeypatch):
    monkeypatch.setattr(pdf_kontoauszug, "ocr_verfuegbar", lambda: False)
    leer = beleg_text(b"", PDF)
    assert leer.quelle == "LEER" and leer.warnungen
    kaputt = beleg_text(b"%PDF-1.4\nkaputt", PDF)
    assert kaputt.quelle == "LEER"
    assert any("nicht lesbar" in w for w in kaputt.warnungen)
    unbekannt = beleg_text(b"\x00\x01\x02", "application/octet-stream", dateiname="x.bin")
    assert unbekannt.quelle == "LEER"
    assert any("nicht ausgelesen" in w for w in unbekannt.warnungen)
    kaputtes_bild = beleg_text(b"\x89PNG\r\n\x1a\nnichts", "image/png")
    assert kaputtes_bild.quelle == "LEER"
    kaputtes_docx = beleg_text(b"PK\x03\x04kaputt", DOCX)
    assert kaputtes_docx.quelle == "LEER"
    assert any("nicht lesbar" in w for w in kaputtes_docx.warnungen)


# ------------------------------------------------------------------ Bilder

def test_bild_ohne_ocr(monkeypatch):
    monkeypatch.setattr(pdf_kontoauszug, "ocr_verfuegbar", lambda: False)
    erg = beleg_text(_png(_seitenbild(ZEILEN)), "image/png")
    assert erg.quelle == "LEER"
    assert erg.seiten == 1
    assert any("tesseract" in w for w in erg.warnungen)


@ocr_noetig
def test_bild_per_ocr():
    bild = _seitenbild(ZEILEN).convert("RGB")
    puffer = BytesIO()
    bild.save(puffer, "JPEG", quality=90)
    erg = beleg_text(puffer.getvalue(), "image/jpeg")
    assert erg.quelle == "OCR", erg.warnungen
    assert "AB-4711" in erg.text


@ocr_noetig
def test_kleines_bild_wird_fuer_ocr_vergroessert():
    erg = beleg_text(_png(_seitenbild(ZEILEN, dpi=100)), "image/png")
    assert erg.quelle == "OCR", erg.warnungen
    assert "4711" in erg.text


# ------------------------------------------------------------------ Office, Mail, Text

def test_docx_mit_kopfzeile_und_tabelle():
    erg = beleg_text(_docx(), DOCX)
    assert erg.quelle == "TEXT"
    assert erg.text.startswith("Nordlicht Handel GmbH · Hafenstraße 5")
    assert erg.text.count("Nordlicht Handel GmbH") == 1
    assert "An: Metallbau Kraft KG" in erg.text
    assert "1 | Stahlträger | 2.500,00" in erg.text
    assert "Summe\t2.500,00 EUR" in erg.text


def test_xlsx_zellwerte():
    erg = beleg_text(_xlsx(), XLSX)
    assert erg.quelle == "TEXT"
    assert erg.seiten == 2
    assert "Tabellenblatt „Bestellung“" in erg.text
    assert "Lieferant | Holz Meier GmbH" in erg.text
    assert "Datum | 2026-11-03" in erg.text
    assert "Eichenbretter | 40 | 12.5" in erg.text
    assert "Summe |  | 500" in erg.text


def test_eml_mit_pdf_anhang():
    erg = beleg_text(_eml(_text_pdf(ZEILEN)), EML)
    assert erg.quelle == "TEXT"
    assert "Betreff: Bestellung 4711 – Büromöbel" in erg.text
    assert "3 Schreibtische" in erg.text
    assert "--- Anhang: bestellung.pdf ---" in erg.text
    assert "AB-4711" in erg.text


def test_eml_nur_html():
    erg = beleg_text(_eml(html=True), "message/rfc822")
    assert "Sehr geehrte Damen und Herren" in erg.text
    assert "Schreibtisch | 899,00" in erg.text
    assert "alert" not in erg.text and "<p>" not in erg.text


def test_txt_kodierungen():
    assert beleg_text("Bestellung für Müller".encode(), "text/plain").text == "Bestellung für Müller"
    erg = beleg_text("Grüße\r\n\r\n\r\n\r\nSumme 10,00 €".encode("cp1252"), "text/plain; charset=windows-1252")
    assert erg.text == "Grüße\n\nSumme 10,00 €"
    assert erg.quelle == "TEXT"
    assert beleg_text(b"   \n ", "text/plain").quelle == "LEER"


def test_mime_nach_inhalt_und_endung():
    assert mime_bestimmen(_text_pdf(ZEILEN), "application/octet-stream") == PDF
    assert mime_bestimmen(_docx(), "application/zip") == DOCX
    assert mime_bestimmen(_xlsx(), None) == XLSX
    assert mime_bestimmen(b"Hallo", "application/octet-stream", "notiz.TXT") == "text/plain"
    assert mime_bestimmen(b"From: x", "", "mail.eml") == EML
    assert mime_bestimmen(b"\xff\xd8\xff\xe0", "image/jpg") == "image/jpeg"
    # Inhalt schlägt Angabe: als PDF deklariertes PNG
    assert mime_bestimmen(_png(_seitenbild(["x"], dpi=30)), "application/pdf") == "image/png"


# ------------------------------------------------------------------ Seitenbilder

def test_vision_pdf_seiten_150_dpi():
    urls = vision_bilder(_text_pdf(ZEILEN, seiten=5), PDF)
    assert len(urls) == 3
    bild = _bild_aus_data_url(urls[0])
    assert bild.format == "PNG"
    assert abs(bild.size[0] - 1240) <= 2 and abs(bild.size[1] - 1754) <= 2   # A4 bei 150 dpi
    assert len(vision_bilder(_text_pdf(ZEILEN, seiten=5), PDF, max_seiten=1)) == 1


def test_vision_nutzt_gemeinsame_pdfium_sperre(monkeypatch):
    original = pdf_kontoauszug._PDFIUM_SPERRE
    eintritte = []

    class Sperre:
        def __enter__(self):
            eintritte.append(1)
            return original.__enter__()

        def __exit__(self, *args):
            return original.__exit__(*args)

    monkeypatch.setattr(pdf_kontoauszug, "_PDFIUM_SPERRE", Sperre())
    assert len(vision_bilder(_text_pdf(ZEILEN), PDF)) == 1
    assert eintritte == [1]


def test_vision_bild_verkleinert_und_ohne_transparenz():
    from PIL import Image

    gross = Image.new("RGBA", (3000, 1000), (255, 0, 0, 0))
    urls = vision_bilder(_png(gross), "image/png")
    assert len(urls) == 1
    bild = _bild_aus_data_url(urls[0])
    assert bild.size == (2000, 667)
    assert bild.mode == "RGB"
    assert bild.getpixel((10, 10)) == (255, 255, 255)   # transparent → weiß


def test_vision_mehrseitiges_tiff():
    from PIL import Image

    seiten = [Image.new("L", (400, 600), farbe) for farbe in (0, 128, 255, 60)]
    puffer = BytesIO()
    seiten[0].save(puffer, "TIFF", save_all=True, append_images=seiten[1:])
    assert len(vision_bilder(puffer.getvalue(), "image/tiff")) == 3
    assert len(vision_bilder(puffer.getvalue(), "image/tiff", max_seiten=2)) == 2


def test_vision_eml_anhang_und_sonstige_typen():
    assert len(vision_bilder(_eml(_text_pdf(ZEILEN, seiten=2)), EML)) == 2
    assert vision_bilder(_eml(), EML) == []
    assert vision_bilder(_docx(), DOCX) == []
    assert vision_bilder(b"%PDF-1.4 kaputt", PDF) == []
    assert vision_bilder(b"", PDF) == []


# ------------------------------------------------------------------ Schutz vor Bomben

from app.services import beleg_text as bt   # noqa: E402

_XNS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
_WNS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _xlsx_roh(blatt: str, geteilt=None) -> bytes:
    """Minimale XLSX von Hand: `blatt` = Inhalt von <sheetData>, `geteilt` = Liste
    von Zeichenketten oder Funktion, die sharedStrings.xml in eine Datei schreibt."""
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ct = "application/vnd.openxmlformats-officedocument.spreadsheetml"
    puffer = BytesIO()
    with zipfile.ZipFile(puffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml",
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   f'<Override PartName="/xl/workbook.xml" ContentType="{ct}.sheet.main+xml"/>'
                   f'<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{ct}.worksheet+xml"/>'
                   f'<Override PartName="/xl/sharedStrings.xml" ContentType="{ct}.sharedStrings+xml"/>'
                   '</Types>')
        z.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   f'<Relationship Id="rId1" Type="{r}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", f'<workbook {_XNS} xmlns:r="{r}"><sheets>'
                   '<sheet name="AB" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels",
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   f'<Relationship Id="rId1" Type="{r}/worksheet" Target="worksheets/sheet1.xml"/>'
                   f'<Relationship Id="rId2" Type="{r}/sharedStrings" Target="sharedStrings.xml"/>'
                   '</Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", f"<worksheet {_XNS}><sheetData>{blatt}</sheetData></worksheet>")
        if callable(geteilt):
            with z.open("xl/sharedStrings.xml", "w", force_zip64=True) as f:
                geteilt(f)
        else:
            z.writestr("xl/sharedStrings.xml", f"<sst {_XNS}>"
                       + "".join(f"<si><t>{s}</t></si>" for s in geteilt or []) + "</sst>")
    return puffer.getvalue()


def _docx_roh(teile: dict[str, str]) -> bytes:
    puffer = BytesIO()
    with zipfile.ZipFile(puffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        for name, inhalt in teile.items():
            z.writestr(name, inhalt)
    return puffer.getvalue()


def _w_dok(inhalt: str) -> str:
    return f"<w:document {_WNS}><w:body>{inhalt}</w:body></w:document>"


def _w_p(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def test_xlsx_archivbombe_wird_abgelehnt():
    def bombe(f):   # rund 30 MB entpackt, wenige KB gepackt
        f.write(f"<sst {_XNS}>".encode())
        block = b"<si><t>x</t></si>" * 10_000
        for _ in range(180):
            f.write(block)
        f.write(b"</sst>")

    daten = _xlsx_roh('<row r="1"><c r="A1" t="s"><v>0</v></c></row>', bombe)
    assert len(daten) < 200_000
    assert mime_bestimmen(daten, None) == XLSX
    erg = beleg_text(daten, XLSX)
    assert erg.quelle == "LEER"
    assert any("entpackt zu groß" in w for w in erg.warnungen), erg.warnungen


def test_xlsx_zu_viele_xml_elemente(monkeypatch):
    monkeypatch.setattr(bt, "ZIP_MAX_ELEMENTE", 1000)
    klein = _xlsx_roh('<row r="1"><c r="A1" t="s"><v>0</v></c></row>', ["Holz Meier"])
    assert beleg_text(klein, XLSX).text.endswith("Holz Meier")
    zeile = "<c><v>1</v></c>" * 2000   # eine Zeile mit vielen Zellen
    erg = beleg_text(_xlsx_roh(f'<row r="1">{zeile}</row>'), XLSX)
    assert erg.quelle == "LEER"
    assert any("XML-Elemente" in w for w in erg.warnungen), erg.warnungen


def test_xlsx_grenzen_je_teil_und_gesamt(monkeypatch):
    monkeypatch.setattr(bt, "XML_MAX_BYTES", 64 * 1024)
    daten = _xlsx_roh('<row r="1"><c r="A1" t="s"><v>0</v></c></row>', ["Stahl " * 20_000])
    erg = beleg_text(daten, XLSX)
    assert erg.quelle == "LEER" and any("entpackt zu groß" in w for w in erg.warnungen)
    monkeypatch.setattr(bt, "XML_MAX_BYTES", 20 * 1024 * 1024)
    monkeypatch.setattr(bt, "ZIP_SUMME_MAX", 64 * 1024)
    assert beleg_text(daten, XLSX).quelle == "LEER"
    monkeypatch.setattr(bt, "ZIP_SUMME_MAX", 50 * 1024 * 1024)
    assert beleg_text(daten, XLSX).quelle == "TEXT"


def test_xlsx_geteilte_zeichenkette_in_allen_zellen(monkeypatch):
    # eine lange geteilte Zeichenkette in 20 × 50 Zellen: nicht Zelle für Zelle vervielfachen
    zeilen = "".join(f'<row r="{z}">' + '<c t="s"><v>0</v></c>' * 50 + "</row>" for z in range(1, 21))
    daten = _xlsx_roh(zeilen, ["A" * 50_000])
    aufrufe = []
    original = bt._zellwert
    monkeypatch.setattr(bt, "_zellwert", lambda w: aufrufe.append(1) or original(w))
    erg = beleg_text(daten, XLSX)
    assert erg.quelle == "TEXT"
    assert len(erg.text) == bt.TEXT_MAX
    assert any("abgeschnitten" in w for w in erg.warnungen)
    assert len(aufrufe) <= 100   # nach der ersten Zeile ist TEXT_MAX erreicht


def test_office_entity_deklaration_abgelehnt():
    dok = ('<?xml version="1.0"?><!DOCTYPE w:document [<!ENTITY a "Bestellung Bestellung">]>'
           + _w_dok(_w_p("&a;&a;&a;")))
    erg = beleg_text(_docx_roh({"word/document.xml": dok}), DOCX)
    assert erg.quelle == "LEER"
    assert "Bestellung" not in erg.text
    assert any("Entity" in w for w in erg.warnungen), erg.warnungen


def test_docx_archivbombe_kopfzeilen():
    puffer = BytesIO()
    with zipfile.ZipFile(puffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", _w_dok(_w_p("Bestellung")))
        for nr in range(1, 3):   # je rund 25 MB entpackt
            with z.open(f"word/header{nr}.xml", "w", force_zip64=True) as f:
                f.write(f"<w:hdr {_WNS}><w:p><w:r><w:t>".encode())
                block = b"y" * 1_000_000
                for _ in range(25):
                    f.write(block)
                f.write(b"</w:t></w:r></w:p></w:hdr>")
    daten = puffer.getvalue()
    assert len(daten) < 200_000
    erg = beleg_text(daten, DOCX)
    assert erg.quelle == "LEER"
    assert any("entpackt zu groß" in w for w in erg.warnungen), erg.warnungen


def test_docx_kopf_und_fusszeilen_begrenzt():
    teile = {"word/document.xml": _w_dok(_w_p("Haupttext"))}
    for nr in range(1, 16):
        teile[f"word/header{nr}.xml"] = f"<w:hdr {_WNS}>{_w_p(f'Kopf {nr} Ende')}</w:hdr>"
        teile[f"word/footer{nr}.xml"] = f"<w:ftr {_WNS}>{_w_p(f'Fuss {nr} Ende')}</w:ftr>"
    erg = beleg_text(_docx_roh(teile), DOCX)
    assert erg.quelle == "TEXT"
    for nr in range(1, 11):   # natürliche Reihenfolge: header2 vor header10
        assert f"Kopf {nr} Ende" in erg.text and f"Fuss {nr} Ende" in erg.text
    for nr in range(11, 16):
        assert f"Kopf {nr} Ende" not in erg.text and f"Fuss {nr} Ende" not in erg.text
    assert "Haupttext" in erg.text


def test_docx_bricht_bei_text_max_ab(monkeypatch):
    monkeypatch.setattr(bt, "TEXT_MAX", 50)
    teile = {
        "word/header1.xml": f"<w:hdr {_WNS}>{_w_p('Briefkopf ' * 10)}</w:hdr>",
        "word/document.xml": _w_dok(_w_p("Haupttext")),
        "word/footer1.xml": "<kaputt",   # wird nach Erreichen von TEXT_MAX nicht mehr gelesen
    }
    erg = beleg_text(_docx_roh(teile), DOCX)
    assert erg.quelle == "TEXT", erg.warnungen
    assert erg.text.startswith("Briefkopf") and len(erg.text) == 50


def test_docx_verschachtelte_tabellen_linear(monkeypatch):
    tief = 60
    xml = _w_p("innen")
    for nr in range(tief):
        xml = f"<w:tbl><w:tr><w:tc>{_w_p(f'z{nr}x')}{xml}</w:tc></w:tr></w:tbl>"
    aufrufe = []
    original = bt._w_absatz
    monkeypatch.setattr(bt, "_w_absatz", lambda p: aufrufe.append(1) or original(p))
    erg = beleg_text(_docx_roh({"word/document.xml": _w_dok(xml)}), DOCX)
    assert erg.quelle == "TEXT"
    assert len(aufrufe) == tief + 1          # jeder Absatz genau einmal (bisher quadratisch)
    for nr in range(tief):
        assert erg.text.count(f"z{nr}x") == 1
    assert erg.text.splitlines()[0] == f"z{tief - 1}x"


def _roh_pdf(seiten: list[tuple[float, float]]) -> bytes:
    """PDF ohne Inhalt mit beliebiger MediaBox je Seite."""
    objekte = ["<< /Type /Catalog /Pages 2 0 R >>",
               "<< /Type /Pages /Kids [" + " ".join(f"{3 + i} 0 R" for i in range(len(seiten)))
               + f"] /Count {len(seiten)} >>"]
    objekte += [f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {b} {h}] >>" for b, h in seiten]
    aus = b"%PDF-1.4\n"
    versatz = []
    for nr, obj in enumerate(objekte, 1):
        versatz.append(len(aus))
        aus += f"{nr} 0 obj\n{obj}\nendobj\n".encode()
    xref = len(aus)
    aus += f"xref\n0 {len(objekte) + 1}\n0000000000 65535 f \n".encode()
    aus += b"".join(f"{v:010d} 00000 n \n".encode() for v in versatz)
    aus += f"trailer\n<< /Size {len(objekte) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return aus


def _render_beobachten(monkeypatch) -> list[tuple[float, bool]]:
    """PdfPage.render mitschreiben: (Maßstab, pdfium-Sperre gehalten)."""
    import pypdfium2 as pdfium

    aufrufe = []
    original = pdfium.PdfPage.render

    def render(self, *args, **kwargs):
        aufrufe.append((kwargs.get("scale"), pdf_kontoauszug._PDFIUM_SPERRE.locked()))
        return original(self, *args, **kwargs)

    monkeypatch.setattr(pdfium.PdfPage, "render", render)
    return aufrufe


def test_vision_uebergrosse_pdf_seite_begrenzt(monkeypatch):
    aufrufe = _render_beobachten(monkeypatch)
    bilder = bt.pdf_seiten_bilder(_roh_pdf([(3000, 2000), (595.28, 841.89)]))
    assert max(bilder[0].size) <= bt.VISION_MAX_KANTE
    assert abs(bilder[0].size[0] - 2000) <= 1 and abs(bilder[0].size[1] - 1333) <= 1
    assert abs(bilder[1].size[0] - 1240) <= 2 and abs(bilder[1].size[1] - 1754) <= 2   # A4: 150 dpi
    assert all(gesperrt for _, gesperrt in aufrufe) and len(aufrufe) == 2
    urls = vision_bilder(_roh_pdf([(14400, 14400)]), PDF)   # größte zulässige MediaBox
    assert len(urls) == 1 and _bild_aus_data_url(urls[0]).size == (2000, 2000)


def test_ocr_uebergrosse_seite_begrenzt_a4_unveraendert(monkeypatch):
    from PIL import Image

    groessen = []

    def tesseract(png, sprache):
        groessen.append(Image.open(BytesIO(png)).size)
        return "level\tpage_num\n"

    monkeypatch.setattr(pdf_kontoauszug, "_tesseract_tsv", tesseract)
    monkeypatch.setattr(pdf_kontoauszug, "_ocr_sprache", lambda: "deu")
    aufrufe = _render_beobachten(monkeypatch)
    seiten = pdf_kontoauszug._ocr_zeilen(_roh_pdf([(595.28, 841.89), (2000, 2000), (842, 595)]))
    assert len(seiten) == 3
    assert abs(groessen[0][0] - 2481) <= 1 and abs(groessen[0][1] - 3508) <= 1   # A4 weiter 300 dpi
    assert abs(groessen[2][0] - 3508) <= 1 and abs(groessen[2][1] - 2481) <= 1   # A4 quer
    breite, hoehe = groessen[1]
    assert breite * hoehe <= pdf_kontoauszug.OCR_MAX_PIXEL * 1.001
    assert breite >= 5990                                # nicht unnötig klein
    assert all(gesperrt for _, gesperrt in aufrufe) and len(aufrufe) == 3


def test_ocr_sperre_auch_bei_fehlern_frei(monkeypatch):
    def kaputt(png, sprache):
        raise RuntimeError("tesseract weg")

    monkeypatch.setattr(pdf_kontoauszug, "_tesseract_tsv", kaputt)
    monkeypatch.setattr(pdf_kontoauszug, "_ocr_sprache", lambda: "deu")
    with pytest.raises(RuntimeError):
        pdf_kontoauszug._ocr_zeilen(_roh_pdf([(595.28, 841.89)]))
    with pytest.raises(Exception):
        pdf_kontoauszug._ocr_zeilen(b"%PDF-1.4 kaputt")
    assert not pdf_kontoauszug._PDFIUM_SPERRE.locked()


def _tiff(seiten: list, **optionen) -> bytes:
    puffer = BytesIO()
    seiten[0].save(puffer, "TIFF", save_all=True, append_images=seiten[1:], **optionen)
    return puffer.getvalue()


def test_bild_dekompressionsbombe_abgelehnt(monkeypatch):
    from PIL import Image

    monkeypatch.setattr(bt, "BILD_MAX_PIXEL", 1_000_000)
    tesseract = []
    monkeypatch.setattr(pdf_kontoauszug, "ocr_verfuegbar", lambda: True)
    monkeypatch.setattr(pdf_kontoauszug, "_tesseract_tsv", lambda png, s: tesseract.append(1) or "")
    seite = Image.new("1", (1500, 1500), 1)   # Fax-TIFF: stark komprimiert, groß dekodiert
    daten = _tiff([seite, seite], compression="group4")
    assert len(daten) < 5000
    erg = beleg_text(daten, "image/tiff")
    assert erg.quelle == "LEER"
    assert any("zu groß" in w for w in erg.warnungen), erg.warnungen
    assert tesseract == []
    assert vision_bilder(daten, "image/tiff") == []


def test_bild_frames_einzeln_dekodiert_und_verkleinert():
    from PIL import Image

    seiten = [Image.new("1", (3000, 1000), 1), Image.new("RGB", (500, 400), "red")]
    frames = bt._bild_frames(_tiff(seiten, compression="tiff_deflate"), 5, max_kante=2000)
    erstes = next(frames)
    assert erstes.size == (2000, 667) and erstes.mode == "L"   # Modus 1 nicht über RGB
    zweites = next(frames)
    assert zweites.size == (500, 400) and zweites.mode == "RGB"
    assert next(frames, None) is None


def test_grosses_jpeg_wird_verkleinert_dekodiert(monkeypatch):
    from PIL import Image

    puffer = BytesIO()
    Image.new("RGB", (4000, 3000), (200, 220, 240)).save(puffer, "JPEG", quality=80)
    monkeypatch.setattr(bt, "BILD_MAX_PIXEL", 5_000_000)   # 12 MP passen nur per JPEG-Draft
    urls = vision_bilder(puffer.getvalue(), "image/jpeg")
    assert len(urls) == 1
    assert _bild_aus_data_url(urls[0]).size == (2000, 1500)


def test_mpo_nur_hauptbild():
    from PIL import Image

    puffer = BytesIO()
    Image.new("RGB", (4000, 3000), "white").save(
        puffer, "MPO", save_all=True, append_images=[Image.new("RGB", (1000, 750), "black")])
    urls = vision_bilder(puffer.getvalue(), "image/jpeg")
    assert len(urls) == 1   # Folgebild ist Tiefen-/HDR-Karte, keine Seite
    assert _bild_aus_data_url(urls[0]).getpixel((5, 5)) == (255, 255, 255)


# ------------------------------------------------------------------ E-Mail: eingebettete Bilder

def _png_bild(groesse, farbe="white") -> bytes:
    from PIL import Image

    return _png(Image.new("RGB", groesse, farbe))


def _eml_mit_signatur(anhaenge: list[tuple[bytes, str, str, str]], logos: int = 3,
                      bild: bytes | None = None) -> bytes:
    """HTML-Mail mit eingebetteten Signaturbildern (multipart/related, cid) und Anhängen."""
    nachricht = EmailMessage()
    nachricht["Subject"] = "Auftragsbestätigung 4711"
    nachricht["From"] = "vertrieb@lieferant.example"
    nachricht["To"] = "einkauf@kunde.example"
    nachricht.set_content("Sehr geehrte Damen und Herren,\nanbei die AB.\n")
    verweise = "".join(f"<img src='cid:logo{i}'>" for i in range(1, logos + 1))
    nachricht.add_alternative(f"<p>anbei die AB.</p>{verweise}", subtype="html")
    html_teil = nachricht.get_payload()[1]
    for i in range(1, logos + 1):
        html_teil.add_related(_png_bild((120, 40), "red"), maintype="image", subtype="png",
                              cid=f"<logo{i}>", filename=f"image00{i}.png", disposition="inline")
    if bild is not None:
        html_teil.add_related(bild, maintype="image", subtype="png", cid="<foto>",
                              filename="foto.png", disposition="inline")
    for inhalt, haupt, unter, name in anhaenge:
        nachricht.add_attachment(inhalt, maintype=haupt, subtype=unter, filename=name)
    return bytes(nachricht)


def test_eml_signaturbilder_verdraengen_pdf_anhang_nicht():
    daten = _eml_mit_signatur([(_text_pdf(ZEILEN, seiten=2), "application", "pdf", "AB_4711.pdf")])
    assert [name for name, _, _ in bt.eml_anhaenge(daten)] == ["AB_4711.pdf"]
    urls = vision_bilder(daten, EML)
    assert len(urls) == 2
    assert all(abs(_bild_aus_data_url(u).size[0] - 1240) <= 2 for u in urls)   # PDF-Seiten
    erg = beleg_text(daten, EML)
    assert "AB-4711" in erg.text
    assert "image001" not in erg.text and not any("image001" in w for w in erg.warnungen)


def test_eml_eingebettetes_bild_ohne_anhang_bleibt():
    foto = _png_bild((900, 700))
    daten = _eml_mit_signatur([], logos=1, bild=foto)
    assert [name for name, _, _ in bt.eml_anhaenge(daten)] == ["foto.png"]   # Logo zu klein
    assert len(vision_bilder(daten, EML)) == 1


def test_vision_eml_pdf_vor_bildern_und_ohne_kleinstbilder():
    daten = _eml_mit_signatur([
        (_png_bild((100, 100)), "image", "png", "symbol.png"),
        (_png_bild((900, 700), "blue"), "image", "png", "foto.png"),
        (_text_pdf(ZEILEN, seiten=2), "application", "pdf", "AB.pdf"),
    ], logos=0)
    urls = vision_bilder(daten, EML)
    groessen = [_bild_aus_data_url(u).size for u in urls]
    assert len(groessen) == 3
    assert all(abs(b - 1240) <= 2 for b, _ in groessen[:2])   # erst die PDF-Seiten
    assert groessen[2] == (900, 700)                           # dann das Foto, kein Symbol
