"""Soll-Referenz, Planstände löschen, Vortrags-Kennzeichen und Exporte."""

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader

from test_api import _login


def _mandant(client, kurz="export"):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": f"{kurz} GmbH", "kurzname": kurz}
    ).json()["id"]
    konten = {k["nummer"]: k for k in
              client.get(f"/api/mandanten/{mandant_id}/konten").json()["konten"]}
    return mandant_id, konten


def _montag(tage_zurueck: int = 0) -> str:
    heute = date.today() - timedelta(days=tage_zurueck)
    return (heute - timedelta(days=heute.weekday())).isoformat()


def test_soll_referenz_persistent_und_loeschbar(client):
    mid, _ = _mandant(client, "referenz")
    alt = client.post(f"/api/mandanten/{mid}/plan/snapshot",
                      json={"start": _montag(14), "kommentar": "Gerichtsplan"}).json()
    neu = client.post(f"/api/mandanten/{mid}/plan/snapshot", json={"start": _montag()}).json()

    # ohne Wahl: jüngster Planstand
    assert client.get(f"/api/mandanten/{mid}/sollist").json()["snapshot"]["id"] == neu["id"]

    # älteren Stand als Referenz wählen -> gilt für Soll/Ist und Plan-Basis
    assert client.put(f"/api/mandanten/{mid}/plan/referenz",
                      json={"snapshot_id": alt["id"]}).status_code == 200
    liste = client.get(f"/api/mandanten/{mid}/plan/snapshots").json()
    assert {s["id"]: s["referenz"] for s in liste} == {alt["id"]: True, neu["id"]: False}
    sollist = client.get(f"/api/mandanten/{mid}/sollist").json()
    assert sollist["snapshot"]["id"] == alt["id"] and sollist["snapshot"]["referenz"]
    plan = client.get(f"/api/mandanten/{mid}/plan").json()
    assert plan["snapshot"]["id"] == alt["id"] and plan["snapshot"]["referenz"]

    # Referenz löschen -> wieder automatisch
    assert client.delete(f"/api/mandanten/{mid}/plan/snapshots/{alt['id']}").status_code == 200
    assert [s["id"] for s in client.get(f"/api/mandanten/{mid}/plan/snapshots").json()] == [neu["id"]]
    assert client.get(f"/api/mandanten/{mid}/sollist").json()["snapshot"]["id"] == neu["id"]

    # zurück auf automatisch, fremde/unbekannte IDs abgewiesen
    assert client.put(f"/api/mandanten/{mid}/plan/referenz", json={"snapshot_id": None}).status_code == 200
    assert client.put(f"/api/mandanten/{mid}/plan/referenz", json={"snapshot_id": 9999}).status_code == 404
    assert client.delete(f"/api/mandanten/{mid}/plan/snapshots/9999").status_code == 404


def test_dauerschuld_vortragen_feld_und_csv(client):
    mid, konten = _mandant(client, "vortragfeld")
    d_id = client.post(f"/api/mandanten/{mid}/dauerbuchungen", json={
        "name": "Miete", "betrag_brutto": 1500, "gueltig_von": "2026-01-01",
        "konto_id": konten["4210"]["id"],
    }).json()["id"]
    liste = client.get(f"/api/mandanten/{mid}/dauerbuchungen").json()
    assert liste[0]["vortragen"] is True                  # Standard: ja
    client.patch(f"/api/dauerbuchungen/{d_id}", json={"vortragen": False})
    assert client.get(f"/api/mandanten/{mid}/dauerbuchungen").json()[0]["vortragen"] is False

    csv = ("Name;Art;Betrag;Intervall;Stichtag;GueltigVon;Konto;Vortrag\n"
           "Leasing;Ausgabe;300;monatlich;5;01.01.2026;4540;nein\n"
           "Telefon;Ausgabe;50;monatlich;10;01.01.2026;4920;\n").encode()
    antwort = client.post(f"/api/mandanten/{mid}/dauerbuchungen/import",
                          files={"datei": ("dauer.csv", csv, "text/csv")})
    assert antwort.json()["angelegt"] == 2
    vortragen = {d["name"]: d["vortragen"]
                 for d in client.get(f"/api/mandanten/{mid}/dauerbuchungen").json()}
    assert vortragen == {"Miete": False, "Leasing": False, "Telefon": True}


def _pdf_seiten(client, mid, **param) -> list[str]:
    antwort = client.get(f"/api/mandanten/{mid}/export/plan.pdf", params=param)
    assert antwort.status_code == 200, antwort.text
    return [s.extract_text() for s in PdfReader(BytesIO(antwort.content)).pages]


def test_pdf_blattaufteilung_nach_spaltenansicht(client):
    mid, _ = _mandant(client, "pdfblatt")
    start = _montag()

    # 26 Wochen: zwei Blätter à 13 Wochen, Kontospalte auf jedem Blatt
    seiten = _pdf_seiten(client, mid, start=start, wochen=26, ansicht="wochen")
    assert len(seiten) == 2
    assert "26-Wochen-Liquiditätsplanung" in seiten[0]
    assert "Blatt 1 von 2" in seiten[0]
    assert "Blatt 2 von 2" in seiten[1]
    for seite in seiten:
        assert "Konto / BWA-Zeile" in seite
        assert seite.count("KW ") <= 13 + 1              # + evtl. Titelzeile

    # 13 Wochen in Wochenansicht: ein Blatt
    assert len(_pdf_seiten(client, mid, start=start, wochen=13, ansicht="wochen")) == 1
    # Tagesansicht: 14 Tage je Blatt -> 13 Wochen = 7 Blätter
    seiten = _pdf_seiten(client, mid, start=start, wochen=13, ansicht="tage")
    assert len(seiten) == 7 and "Tageswerte" in seiten[-1]
    # Monatsansicht: ein Blatt mit Monatsnamen
    seiten = _pdf_seiten(client, mid, start=start, wochen=26, ansicht="monate", metrik="vortrag")
    assert len(seiten) == 1 and "Monatswerte" in seiten[0] and "mit Vortrag" in seiten[0]

    assert client.get(f"/api/mandanten/{mid}/export/plan.pdf",
                      params={"ansicht": "quartale"}).status_code == 422
    assert client.get(f"/api/mandanten/{mid}/export/plan.pdf",
                      params={"metrik": "egal"}).status_code == 422


def test_excel_monatsblatt_und_referenz(client):
    mid, _ = _mandant(client, "xlsxref")
    snap = client.post(f"/api/mandanten/{mid}/plan/snapshot", json={"start": _montag()}).json()
    client.put(f"/api/mandanten/{mid}/plan/referenz", json={"snapshot_id": snap["id"]})
    antwort = client.get(f"/api/mandanten/{mid}/export/plan.xlsx",
                         params={"wochen": 13, "metrik": "vortrag"})
    assert antwort.status_code == 200
    wb = load_workbook(BytesIO(antwort.content))
    assert wb.sheetnames == ["Info", "Tage", "Wochen", "Monate", "Soll-Ist"]
    monate = wb["Monate"]
    assert monate.cell(row=2, column=2).value == "Plan"
    assert monate.cell(row=1, column=2).value.split()[0] in (
        "Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August",
        "September", "Oktober", "November", "Dezember")
    assert "gewählte Soll-Referenz" in wb["Soll-Ist"].cell(row=1, column=1).value
    info = {wb["Info"].cell(row=i, column=1).value: wb["Info"].cell(row=i, column=2).value
            for i in range(1, 20)}
    assert info["Soll-Basis"].startswith("gewählte Soll-Referenz")
    assert "mit Vortrag" in info["Werte"]
