def _login(client):
    antwort = client.post(
        "/login", data={"email": "admin@example.com", "passwort": "admin"},
        follow_redirects=False,
    )
    assert antwort.status_code == 303


def test_login_und_mandant_workflow(client):
    # ohne Login: API 401, Seite -> Redirect
    assert client.get("/api/mandanten").status_code == 401
    assert client.get("/", follow_redirects=False).status_code == 303

    _login(client)
    assert client.get("/api/mandanten").json() == []

    antwort = client.post(
        "/api/mandanten",
        json={"name": "Testfirma GmbH", "kurzname": "testfirma", "kontenrahmen": "SKR04",
              "bundesland": "BY", "verfahrensstatus": "EROEFFNET"},
    )
    assert antwort.status_code == 201, antwort.text
    mandant_id = antwort.json()["id"]

    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()
    nummern = {k["nummer"] for k in konten["konten"]}
    assert {"1800", "4400", "6110", "3820"} <= nummern
    k4400 = next(k for k in konten["konten"] if k["nummer"] == "4400")
    assert k4400["ust_satz"] == 19.0

    plan = client.get(f"/api/mandanten/{mandant_id}/plan").json()
    assert len(plan["tage"]) == 13 * 7
    assert plan["vergleichsbasis"] == "LIVE"

    # Termine generieren und im Plan sehen
    regeln = client.get(f"/api/mandanten/{mandant_id}/termin-regeln").json()
    sv = next(r for r in regeln if r["typ"] == "SV")
    client.patch(f"/api/termin-regeln/{sv['id']}", json={"betrag_modus": "FIX", "betrag_fix": 5000})
    erg = client.post(f"/api/mandanten/{mandant_id}/termine/generieren", json={}).json()
    assert erg["angelegt"] >= 3

    snapshot = client.post(f"/api/mandanten/{mandant_id}/plan/snapshot", json={})
    assert snapshot.status_code == 201
    sollist = client.get(f"/api/mandanten/{mandant_id}/sollist").json()
    assert sollist["snapshot"]["id"] == snapshot.json()["id"]


def test_import_datev(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Import AG", "kurzname": "import-ag"}
    ).json()["id"]

    extf = (
        '"EXTF";700;21;"Buchungsstapel";12;;;"";"";"";1;1;20260101;4;20260801;20260831;"";"";1\n'
        "Umsatz (ohne Soll/Haben-Kz);Soll/Haben-Kennzeichen;WKZ;Kurs;Basis;WKZ2;Konto;"
        "Gegenkonto (ohne BU-Schlüssel);BU;Belegdatum;Belegfeld 1;Belegfeld 2;Skonto;Buchungstext\n"
        '1.190,00;"S";;;;;1200;8400;;0308;"RE-1";;;"Zahlungseingang"\n'
    ).encode("cp1252")
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/import",
        files={"datei": ("buchungen.csv", extf, "text/csv")},
    )
    assert antwort.status_code == 200, antwort.text
    erg = antwort.json()
    assert erg["format"] == "DATEV"
    assert erg["anzahl"] == 1


def test_mandantentrennung(client):
    _login(client)
    m1 = client.post("/api/mandanten", json={"name": "A", "kurzname": "a"}).json()["id"]

    # Bearbeiter ohne Mandantenzuordnung anlegen
    client.post(
        "/api/benutzer",
        json={"email": "b@example.com", "name": "B", "passwort": "geheim123",
              "rolle": "BEARBEITER", "mandanten_ids": []},
    )
    client.get("/logout")
    client.post("/login", data={"email": "b@example.com", "passwort": "geheim123"})
    assert client.get("/api/mandanten").json() == []
    assert client.get(f"/api/mandanten/{m1}/plan").status_code == 403
    # keine Admin-Funktionen
    assert client.get("/api/benutzer").status_code == 403


def test_forderungsklasse_automatik(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten",
        json={"name": "InsO GmbH", "kurzname": "inso", "verfahrensstatus": "VORLAEUFIG",
              "insolvenz_stichtag": "2026-07-15"},
    ).json()["id"]

    # Rechnungsdatum vor dem Stichtag -> Insolvenzforderung
    alt = client.post(
        f"/api/mandanten/{mandant_id}/posten",
        json={"art": "KREDITOR", "partner": "Alt", "rechnungsdatum": "2026-07-01",
              "faellig_am": "2026-08-20", "betrag_brutto": 5000},
    ).json()
    assert alt["forderungsklasse"] == "INSOLVENZFORDERUNG"
    # nach dem Stichtag -> Masse
    neu = client.post(
        f"/api/mandanten/{mandant_id}/posten",
        json={"art": "KREDITOR", "partner": "Neu", "rechnungsdatum": "2026-07-20",
              "faellig_am": "2026-08-20", "betrag_brutto": 700},
    ).json()
    assert neu["forderungsklasse"] == "MASSE"
    # explizite Angabe hat Vorrang; ungültige Klasse wird abgelehnt
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/posten",
        json={"art": "KREDITOR", "partner": "X", "rechnungsdatum": "2026-07-01",
              "faellig_am": "2026-08-20", "betrag_brutto": 1,
              "forderungsklasse": "QUATSCH"},
    )
    assert antwort.status_code == 422
    # gesperrte Forderung erscheint nicht im Plan, aber im Sperr-Ausweis
    plan = client.get(f"/api/mandanten/{mandant_id}/plan").json()
    assert plan["gesperrte_insolvenzforderungen"] == 5000.0


def test_export_xlsx_und_pdf(client):
    from io import BytesIO

    from openpyxl import load_workbook

    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Export AG", "kurzname": "export-ag"}
    ).json()["id"]
    client.post(f"/api/mandanten/{mandant_id}/plan/snapshot", json={})

    antwort = client.get(f"/api/mandanten/{mandant_id}/export/plan.xlsx")
    assert antwort.status_code == 200
    assert "spreadsheetml" in antwort.headers["content-type"]
    wb = load_workbook(BytesIO(antwort.content))
    assert {"Info", "Tage", "Wochen", "Soll-Ist"} <= set(wb.sheetnames)
    assert wb["Info"]["B2"].value == "Export AG"

    antwort = client.get(f"/api/mandanten/{mandant_id}/export/plan.pdf")
    assert antwort.status_code == 200
    assert antwort.headers["content-type"] == "application/pdf"
    assert antwort.content.startswith(b"%PDF")

    # Export im Szenario: läuft durch und Kopfzeilen-Helfer liefern den Vermerk
    szenarien = client.get(f"/api/mandanten/{mandant_id}/szenarien").json()
    worst = next(s for s in szenarien if s["name"] == "Worst Case")
    antwort = client.get(
        f"/api/mandanten/{mandant_id}/export/plan.pdf?szenario_id={worst['id']}"
    )
    assert antwort.status_code == 200
    plan = client.get(
        f"/api/mandanten/{mandant_id}/plan?szenario_id={worst['id']}"
    ).json()
    from app.services.export import _szenario_text

    assert "Worst Case" in _szenario_text(plan)
    assert "Debitoren +14 Tage" in _szenario_text(plan)


def test_szenarien_crud_und_plan(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Szenario GmbH", "kurzname": "szen"}
    ).json()["id"]

    # Standard-Vorlagen wurden angelegt
    szenarien = client.get(f"/api/mandanten/{mandant_id}/szenarien").json()
    namen = {s["name"] for s in szenarien}
    assert {"Best Case", "Worst Case"} <= namen
    worst = next(s for s in szenarien if s["name"] == "Worst Case")

    plan = client.get(
        f"/api/mandanten/{mandant_id}/plan?szenario_id={worst['id']}"
    ).json()
    assert plan["szenario"]["name"] == "Worst Case"

    # CRUD + Validierung
    neu = client.post(
        f"/api/mandanten/{mandant_id}/szenarien",
        json={"name": "Stress", "ein_faktor": 60, "aus_faktor": 100,
              "debitoren_verzoegerung_tage": 30},
    )
    assert neu.status_code == 201
    assert client.patch(
        f"/api/szenarien/{neu.json()['id']}", json={"ein_faktor": 9999}
    ).status_code == 422
    assert client.delete(f"/api/szenarien/{neu.json()['id']}").json()["ok"] is True
    # fremdes/unbekanntes Szenario
    assert client.get(
        f"/api/mandanten/{mandant_id}/plan?szenario_id=99999"
    ).status_code == 404


def test_insolvenzgeld_einstellungen_und_vorschau(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten",
        json={"name": "IG GmbH", "kurzname": "ig", "verfahrensstatus": "VORLAEUFIG",
              "insolvenz_stichtag": "2026-07-15"},
    ).json()["id"]
    antwort = client.patch(
        f"/api/mandanten/{mandant_id}",
        json={"insolvenzgeld_aktiv": True, "insolvenzgeld_von": "2026-06-10",
              "insolvenzgeld_bis": "2026-09-09"},
    )
    assert antwort.status_code == 200
    vorschau = client.get(
        f"/api/mandanten/{mandant_id}/insolvenzgeld/vorschau"
    ).json()
    assert vorschau["aktiv"] is True
    assert vorschau["von"] == "2026-06-10"
    plan = client.get(f"/api/mandanten/{mandant_id}/plan").json()
    assert plan["insolvenzgeld"]["aktiv"] is True


MT940_BEISPIEL = (
    ":20:STMT-1\n"
    ":25:DE89370400440532013000\n"
    ":28C:152/1\n"
    ":60F:C260803EUR41250,00\n"
    ":61:2608040804CR1190,00NTRFNONREF\n"
    ":86:166?00GUTSCHRIFT?20SVWZ+RE-100?32Kunde Albrecht AG\n"
    ":61:260805DR500,00NDDT0815\n"
    ":86:105?20SVWZ+Miete?32Vermieter GmbH\n"
    ":62F:C260805EUR41940,00\n"
    "-\n"
).encode("cp1252")


def test_bank_import_mt940(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Bank GmbH", "kurzname": "bank"}
    ).json()["id"]
    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()["konten"]
    bank_konto = next(k for k in konten if k["nummer"] == "1200")
    # IBAN mit Leerzeichen -> wird normalisiert gespeichert
    client.patch(f"/api/konten/{bank_konto['id']}", json={"iban": "DE89 3704 0044 0532 0130 00"})

    antwort = client.post(
        f"/api/mandanten/{mandant_id}/bank-import",
        files={"datei": ("auszug.sta", MT940_BEISPIEL, "text/plain")},
    )
    assert antwort.status_code == 200, antwort.text
    erg = antwort.json()
    assert erg["format"] == "MT940"
    assert erg["anzahl"] == 2
    assert erg["bestandsanker"][0]["datum"] == "2026-08-05"
    assert erg["bestandsanker"][0]["wert"] == 41940.0

    umsaetze = client.get(f"/api/mandanten/{mandant_id}/bank-umsaetze").json()
    assert len(umsaetze) == 2
    assert umsaetze[0]["partner"] == "Vermieter GmbH"  # neuester zuerst

    # Endsaldo wirkt als Bestandsanker in der Planung
    plan = client.get(f"/api/mandanten/{mandant_id}/plan").json()
    bank_zeile = next(
        f for f in plan["bestaende"]["finanzkonten"] if f["nummer"] == "1200"
    )
    assert bank_zeile["anker"] == {"datum": "2026-08-05", "wert": 41940.0}
    assert plan["bestaende"]["liquiditaet"][plan["heute"]] == 41940.0

    # Re-Import: Dubletten-Warnung; Anker wird aktualisiert, nicht dupliziert
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/bank-import",
        files={"datei": ("auszug.sta", MT940_BEISPIEL, "text/plain")},
    )
    assert any("Dubletten" in w for w in antwort.json()["warnungen"])
    bestaende = client.get(f"/api/mandanten/{mandant_id}/bestaende").json()
    assert len([b for b in bestaende if b["datum"] == "2026-08-05"]) == 1

    # Import löschen entfernt die Bankumsätze des Batches
    client.delete(f"/api/importe/{erg['batch_id']}")
    umsaetze = client.get(f"/api/mandanten/{mandant_id}/bank-umsaetze").json()
    assert len(umsaetze) == 2  # die des zweiten Imports bleiben


def test_bank_import_ohne_zuordnung(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Bank2 GmbH", "kurzname": "bank2"}
    ).json()["id"]
    # keine IBAN hinterlegt, aber zwei Finanzkonten (Bank + Kasse) -> Warnung, 0 Umsätze
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/bank-import",
        files={"datei": ("auszug.sta", MT940_BEISPIEL, "text/plain")},
    )
    erg = antwort.json()
    assert erg["anzahl"] == 0
    assert any("übersprungen" in w for w in erg["warnungen"])
    # mit expliziter Kontoauswahl klappt es
    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()["konten"]
    bank_konto = next(k for k in konten if k["nummer"] == "1200")
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/bank-import?konto_id={bank_konto['id']}",
        files={"datei": ("auszug.sta", MT940_BEISPIEL, "text/plain")},
    )
    assert antwort.json()["anzahl"] == 2


def test_op_ausgleich_flow(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Abgleich GmbH", "kurzname": "abgleich"}
    ).json()["id"]
    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()["konten"]
    bank_konto = next(k for k in konten if k["nummer"] == "1200")
    client.patch(f"/api/konten/{bank_konto['id']}", json={"iban": "DE89370400440532013000"})
    client.post(
        f"/api/mandanten/{mandant_id}/bank-import",
        files={"datei": ("auszug.sta", MT940_BEISPIEL, "text/plain")},
    )
    # passende offene Posten anlegen
    fo = client.post(
        f"/api/mandanten/{mandant_id}/posten",
        json={"art": "DEBITOR", "partner": "Kunde Albrecht AG", "belegnr": "RE-100",
              "faellig_am": "2026-08-10", "betrag_brutto": 1190},
    ).json()["id"]
    er = client.post(
        f"/api/mandanten/{mandant_id}/posten",
        json={"art": "KREDITOR", "partner": "Vermieter GmbH",
              "faellig_am": "2026-08-05", "betrag_brutto": 500},
    ).json()["id"]

    vorschlaege = client.get(
        f"/api/mandanten/{mandant_id}/op-ausgleich/vorschlaege"
    ).json()
    assert len(vorschlaege) == 2
    assert all(v["konfidenz"] == "SICHER" for v in vorschlaege)

    paare = [
        {"transaktion_id": v["transaktion"]["id"], "posten_id": v["posten"]["id"]}
        for v in vorschlaege
    ]
    erg = client.post(
        f"/api/mandanten/{mandant_id}/op-ausgleich", json={"paare": paare}
    ).json()
    assert erg["ausgeglichen"] == 2
    posten = client.get(f"/api/mandanten/{mandant_id}/posten").json()
    assert all(p["status"] == "BEZAHLT" for p in posten)
    assert next(p for p in posten if p["id"] == fo)["bezahlt_am"] == "2026-08-04"
    # keine neuen Vorschläge mehr; erneuter Ausgleich desselben Umsatzes -> 409
    assert client.get(f"/api/mandanten/{mandant_id}/op-ausgleich/vorschlaege").json() == []
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/op-ausgleich", json={"paare": [paare[0]]}
    )
    assert antwort.status_code == 409

    # Aufheben öffnet den Posten wieder
    client.post(
        f"/api/mandanten/{mandant_id}/op-ausgleich/aufheben",
        json={"transaktion_id": paare[0]["transaktion_id"]},
    )
    posten = client.get(f"/api/mandanten/{mandant_id}/posten").json()
    assert sum(1 for p in posten if p["status"] == "OFFEN") == 1
    assert len(client.get(f"/api/mandanten/{mandant_id}/op-ausgleich/vorschlaege").json()) == 1


def test_szenarien_vergleich_endpoint(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Vergleich AG", "kurzname": "vgl"}
    ).json()["id"]
    v = client.get(f"/api/mandanten/{mandant_id}/szenarien-vergleich").json()
    namen = [s["name"] for s in v["szenarien"]]
    assert namen[0] == "Basisplan"
    assert {"Best Case", "Worst Case"} <= set(namen)
    assert len(v["szenarien"][0]["liquiditaet_wochenende"]) == len(v["wochen"])
    # Seite rendert
    assert client.get(f"/mandanten/{mandant_id}/szenarien").status_code == 200


def test_zwei_faktor_anmeldung(client):
    import time

    from app.services.totp import code_fuer_schritt

    _login(client)
    # Einrichten und mit gültigem Code bestätigen
    d = client.post("/api/profil/2fa/einrichten").json()
    assert d["qr_svg"].startswith("data:image/svg+xml")
    schritt = int(time.time() // 30)
    antwort = client.post("/api/profil/2fa/bestaetigen",
                          json={"code": code_fuer_schritt(d["geheimnis"], schritt)})
    assert antwort.status_code == 200, antwort.text
    assert client.get("/api/profil").json()["totp_aktiv"] is True

    # Neuer Login: Passwort allein reicht nicht mehr
    client.get("/logout")
    antwort = client.post(
        "/login", data={"email": "admin@example.com", "passwort": "admin"},
        follow_redirects=False,
    )
    assert antwort.headers["location"] == "/login/2fa"
    assert client.get("/api/mandanten").status_code == 401  # noch keine Sitzung

    # falscher Code -> abgelehnt; Folgeschritt-Code -> angemeldet
    antwort = client.post("/login/2fa", data={"code": "000000"}, follow_redirects=False)
    assert antwort.status_code == 401
    code = code_fuer_schritt(d["geheimnis"], schritt + 1)
    antwort = client.post("/login/2fa", data={"code": code}, follow_redirects=False)
    assert antwort.status_code == 303
    assert client.get("/api/mandanten").status_code == 200

    # Deaktivieren mit Passwort
    antwort = client.post("/api/profil/2fa/deaktivieren", json={"passwort": "admin"})
    assert antwort.status_code == 200
    assert client.get("/api/profil").json()["totp_aktiv"] is False


def test_eigenes_passwort_aendern(client):
    _login(client)
    assert client.post(
        "/api/profil/passwort",
        json={"aktuelles_passwort": "falsch", "neues_passwort": "neues-passwort"},
    ).status_code == 403
    assert client.post(
        "/api/profil/passwort",
        json={"aktuelles_passwort": "admin", "neues_passwort": "neues-passwort"},
    ).status_code == 200
    client.get("/logout")
    antwort = client.post(
        "/login", data={"email": "admin@example.com", "passwort": "neues-passwort"},
        follow_redirects=False,
    )
    assert antwort.status_code == 303


def test_szenario_detailregeln_api(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Regel AG", "kurzname": "regel-ag"}
    ).json()["id"]
    szenarien = client.get(f"/api/mandanten/{mandant_id}/szenarien").json()
    worst = next(s for s in szenarien if s["name"] == "Worst Case")
    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()
    k8400 = next(k for k in konten["konten"] if k["nummer"] == "8400")
    gruppe = next(g for g in konten["gruppen"] if g["code"] == "E_UMSATZ")

    # beides gesetzt -> 422
    antwort = client.put(
        f"/api/szenarien/{worst['id']}/regeln",
        json=[{"konto_id": k8400["id"], "gruppe_id": gruppe["id"], "faktor": 50}],
    )
    assert antwort.status_code == 422
    antwort = client.put(
        f"/api/szenarien/{worst['id']}/regeln",
        json=[{"konto_id": k8400["id"], "faktor": 50},
              {"gruppe_id": gruppe["id"], "faktor": 60}],
    )
    assert antwort.status_code == 200
    regeln = client.get(f"/api/szenarien/{worst['id']}/regeln").json()
    assert len(regeln) == 2
    plan = client.get(
        f"/api/mandanten/{mandant_id}/plan?szenario_id={worst['id']}"
    ).json()
    assert plan["szenario"]["regeln_anzahl"] == 2


def test_leser_darf_nicht_schreiben(client):
    _login(client)
    m1 = client.post("/api/mandanten", json={"name": "C", "kurzname": "c"}).json()["id"]
    client.post(
        "/api/benutzer",
        json={"email": "l@example.com", "name": "L", "passwort": "geheim123",
              "rolle": "LESER", "mandanten_ids": [m1]},
    )
    client.get("/logout")
    client.post("/login", data={"email": "l@example.com", "passwort": "geheim123"})
    assert client.get(f"/api/mandanten/{m1}/plan").status_code == 200
    antwort = client.post(
        f"/api/mandanten/{m1}/posten",
        json={"art": "KREDITOR", "partner": "X", "faellig_am": "2026-09-01",
              "betrag_brutto": 100},
    )
    assert antwort.status_code == 403


def test_op_import(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten",
        json={"name": "OP Import GmbH", "kurzname": "op-import",
              "verfahrensstatus": "VORLAEUFIG", "insolvenz_stichtag": "2026-07-15"},
    ).json()["id"]

    csv_daten = (
        "Art;Partner;Belegnummer;Rechnungsdatum;Faellig;Betrag;Konto;Notiz\n"
        "ER;Altlieferant Nord;RE-ALT-1;01.06.2026;01.09.2026;4000,00;3400;vor Stichtag\n"
        "ER;Neulieferant Sued;RE-NEU-1;01.08.2026;01.09.2026;2000,00;3400;nach Stichtag\n"
        "FO;Kunde West;RE-2026-9;20.07.2026;19.08.2026;11900,00;8400;\n"
        "ER;Unbekanntes Konto AG;RE-U-1;;15.09.2026;500,00;9999;\n"
    ).encode("utf-8")
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/posten/import",
        files={"datei": ("op-liste.csv", csv_daten, "text/csv")},
    )
    assert antwort.status_code == 200, antwort.text
    erg = antwort.json()
    assert erg["angelegt"] == 4
    assert erg["uebersprungen"] == 0
    assert erg["insolvenzforderungen"] == 1
    assert any("9999" in w for w in erg["warnungen"])

    posten = client.get(f"/api/mandanten/{mandant_id}/posten").json()
    nach_beleg = {p["belegnr"]: p for p in posten}
    assert nach_beleg["RE-ALT-1"]["forderungsklasse"] == "INSOLVENZFORDERUNG"
    assert nach_beleg["RE-NEU-1"]["forderungsklasse"] == "MASSE"
    assert nach_beleg["RE-2026-9"]["forderungsklasse"] is None
    assert nach_beleg["RE-U-1"]["konto_id"] is None

    # erneuter Import: alles Duplikate
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/posten/import",
        files={"datei": ("op-liste.csv", csv_daten, "text/csv")},
    )
    erg = antwort.json()
    assert erg["angelegt"] == 0
    assert erg["uebersprungen"] == 4


def test_kontenrahmen_vorlagen(client):
    _login(client)
    liste = client.get("/api/kontenrahmen").json()
    assert {v["code"] for v in liste} >= {"SKR03", "SKR04", "EIGEN"}

    antwort = client.post(
        "/api/kontenrahmen",
        json={"code": "SKR03-INSO", "name": "SKR03 Insolvenz", "basis": "SKR03"},
    )
    assert antwort.status_code == 201, antwort.text
    definition = client.get("/api/kontenrahmen/SKR03-INSO").json()
    assert definition["fest"] is False
    assert len(definition["konten"]) > 40          # Kopie der SKR03-Vorlage

    # Vorlage erweitern: eigene Gruppe + Konto mit Gruppenzuordnung
    definition["gruppen"].append(
        {"code": "A_MASSE", "name": "Massekosten", "richtung": "AUS", "sortierung": 125}
    )
    definition["konten"].append(
        {"nummer": "4999", "bezeichnung": "Massekosten Verfahren", "typ": "AUFWAND",
         "ust_satz": 19.0, "gruppe": "A_MASSE"}
    )
    assert client.put("/api/kontenrahmen/SKR03-INSO", json=definition).status_code == 200

    # Mandant aus der eigenen Vorlage
    mandant_id = client.post(
        "/api/mandanten",
        json={"name": "Vorlagen GmbH", "kurzname": "vorlagen", "kontenrahmen": "SKR03-INSO"},
    ).json()["id"]
    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()
    gruppe = next(g for g in konten["gruppen"] if g["code"] == "A_MASSE")
    k4999 = next(k for k in konten["konten"] if k["nummer"] == "4999")
    assert k4999["gruppe_id"] == gruppe["id"]
    assert k4999["ust_satz"] == 19.0

    # eingebaute Vorlagen sind geschützt, Unbekanntes wird abgelehnt
    assert client.put("/api/kontenrahmen/SKR03", json=definition).status_code == 409
    assert client.delete("/api/kontenrahmen/SKR03").status_code == 409
    assert client.post(
        "/api/mandanten",
        json={"name": "X", "kurzname": "x-kr", "kontenrahmen": "GIBTSNICHT"},
    ).status_code == 422
    kaputt = {"gruppen": [], "konten": [
        {"nummer": "100", "bezeichnung": "A", "typ": "AUFWAND"},
        {"nummer": "100", "bezeichnung": "B", "typ": "AUFWAND"},
    ]}
    assert client.put("/api/kontenrahmen/SKR03-INSO", json=kaputt).status_code == 422

    assert client.delete("/api/kontenrahmen/SKR03-INSO").json()["ok"] is True
    # Vorlagenverwaltung ist Admin-Sache
    client.post("/api/benutzer", json={
        "email": "kr@example.com", "name": "K", "passwort": "geheim123",
        "rolle": "BEARBEITER", "mandanten_ids": []})
    client.get("/logout")
    client.post("/login", data={"email": "kr@example.com", "passwort": "geheim123"})
    assert client.get("/api/kontenrahmen").status_code == 403


def test_gruppen_crud(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Gruppen GmbH", "kurzname": "gruppen"}
    ).json()["id"]
    gid = client.post(
        f"/api/mandanten/{mandant_id}/gruppen",
        json={"code": "A_MASSE", "name": "Massekosten", "richtung": "AUS", "sortierung": 125},
    ).json()["id"]
    assert client.post(
        f"/api/mandanten/{mandant_id}/gruppen",
        json={"code": "A_MASSE", "name": "Doppelt"},
    ).status_code == 409
    assert client.patch(
        f"/api/gruppen/{gid}", json={"name": "Masseverbindlichkeiten", "sortierung": 10}
    ).json()["ok"] is True

    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()
    gruppe = next(g for g in konten["gruppen"] if g["id"] == gid)
    assert gruppe["name"] == "Masseverbindlichkeiten"
    assert gruppe["sortierung"] == 10

    # Konto zuordnen, Gruppe löschen -> Konto bleibt ohne Gruppe
    k4210 = next(k for k in konten["konten"] if k["nummer"] == "4210")
    client.patch(f"/api/konten/{k4210['id']}", json={"gruppe_id": gid})
    assert client.delete(f"/api/gruppen/{gid}").json()["ok"] is True
    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()
    assert gid not in {g["id"] for g in konten["gruppen"]}
    assert next(k for k in konten["konten"] if k["nummer"] == "4210")["gruppe_id"] is None


def test_zahlung_erfassen(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Zahlung GmbH", "kurzname": "zahlung"}
    ).json()["id"]
    pid = client.post(
        f"/api/mandanten/{mandant_id}/posten",
        json={"art": "DEBITOR", "partner": "Kunde Nord", "faellig_am": "2026-11-20",
              "betrag_brutto": 10000},
    ).json()["id"]

    # Teilzahlung: Posten bleibt offen, Planung rechnet mit dem Rest
    erg = client.post(f"/api/posten/{pid}/zahlungen",
                      json={"betrag": 4000, "datum": "2026-11-18"}).json()
    assert erg["status"] == "OFFEN"
    assert erg["restbetrag"] == 6000.0
    # Überzahlung auf normalem Posten wird abgelehnt
    assert client.post(f"/api/posten/{pid}/zahlungen",
                       json={"betrag": 9000}).status_code == 422
    # Skonto: Rest ausbuchen
    erg = client.post(f"/api/posten/{pid}/zahlungen",
                      json={"betrag": 5800, "ausbuchen": True}).json()
    assert erg["status"] == "BEZAHLT"
    posten = client.get(f"/api/mandanten/{mandant_id}/posten").json()
    assert posten[0]["status"] == "BEZAHLT"
    assert posten[0]["bezahlt_betrag"] == 9800.0
    # Zahlung auf bezahlten Posten abgelehnt
    assert client.post(f"/api/posten/{pid}/zahlungen",
                       json={"betrag": 1}).status_code == 422


def test_sammelposten_ueberzahlung_vortrag(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Sammel GmbH", "kurzname": "sammel"}
    ).json()["id"]
    wochen = [("Debitoren KW 46", "2026-11-13", 10000),
              ("Debitoren KW 47", "2026-11-20", 8000),
              ("Debitoren KW 48", "2026-11-27", 6000)]
    ids = []
    for partner, faellig, betrag in wochen:
        ids.append(client.post(
            f"/api/mandanten/{mandant_id}/posten",
            json={"art": "DEBITOR", "partner": partner, "faellig_am": faellig,
                  "betrag_brutto": betrag, "sammel": True},
        ).json()["id"])

    # Mehr eingezahlt als erwartet: 15.000 auf KW 46 -> 5.000 auf KW 47 vorgetragen
    erg = client.post(f"/api/posten/{ids[0]}/zahlungen", json={"betrag": 15000}).json()
    assert erg["status"] == "BEZAHLT"
    assert erg["verrechnet"] == [{
        "posten_id": ids[1], "partner": "Debitoren KW 47", "faellig_am": "2026-11-20",
        "betrag": 5000.0, "status": "OFFEN"}]
    assert erg["unverteilt"] == 0.0
    posten = {p["partner"]: p for p in client.get(f"/api/mandanten/{mandant_id}/posten").json()}
    assert posten["Debitoren KW 47"]["restbetrag"] == 3000.0    # spätere Einzahlung sinkt

    # 12.000 auf KW 47 (Rest 3.000): KW 47 und KW 48 bezahlt, 3.000 nicht verteilbar
    erg = client.post(f"/api/posten/{ids[1]}/zahlungen", json={"betrag": 12000}).json()
    assert erg["status"] == "BEZAHLT"
    assert [v["posten_id"] for v in erg["verrechnet"]] == [ids[2]]
    assert erg["verrechnet"][0]["betrag"] == 6000.0
    assert erg["verrechnet"][0]["status"] == "BEZAHLT"
    assert erg["unverteilt"] == 3000.0


def test_mandant_loeschen(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten",
        json={"name": "Wegfall GmbH", "kurzname": "wegfall",
              "verfahrensstatus": "VORLAEUFIG", "insolvenz_stichtag": "2026-07-15"},
    ).json()["id"]

    # Daten über möglichst viele Tabellen verteilen
    konten = client.get(f"/api/mandanten/{mandant_id}/konten").json()["konten"]
    k1200 = next(k for k in konten if k["nummer"] == "1200")
    k8400 = next(k for k in konten if k["nummer"] == "8400")
    client.patch(f"/api/konten/{k1200['id']}", json={"iban": "DE89370400440532013000"})
    client.post(f"/api/mandanten/{mandant_id}/posten", json={
        "art": "DEBITOR", "partner": "Kunde Albrecht AG", "belegnr": "RE-100",
        "faellig_am": "2026-08-10", "betrag_brutto": 1190, "konto_id": k8400["id"]})
    client.post(f"/api/mandanten/{mandant_id}/bank-import",
                files={"datei": ("auszug.sta", MT940_BEISPIEL, "text/plain")})
    vorschlaege = client.get(f"/api/mandanten/{mandant_id}/op-ausgleich/vorschlaege").json()
    client.post(f"/api/mandanten/{mandant_id}/op-ausgleich", json={"paare": [{
        "transaktion_id": vorschlaege[0]["transaktion"]["id"],
        "posten_ids": [vorschlaege[0]["posten"]["id"]]}]})
    client.put(f"/api/mandanten/{mandant_id}/budgets", json=[
        {"konto_id": k8400["id"], "jahr": 2026, "monat": 9, "betrag_netto": 50000}])
    regeln = client.get(f"/api/mandanten/{mandant_id}/termin-regeln").json()
    client.patch(f"/api/termin-regeln/{regeln[0]['id']}",
                 json={"betrag_modus": "FIX", "betrag_fix": 5000})
    client.post(f"/api/mandanten/{mandant_id}/termine/generieren", json={})
    client.post(f"/api/mandanten/{mandant_id}/plan/snapshot", json={})
    szenarien = client.get(f"/api/mandanten/{mandant_id}/szenarien").json()
    client.put(f"/api/szenarien/{szenarien[0]['id']}/regeln",
               json=[{"konto_id": k8400["id"], "faktor": 80}])

    # Bestätigung erforderlich; Bearbeiter dürfen gar nicht löschen
    assert client.delete(f"/api/mandanten/{mandant_id}").status_code == 422
    assert client.delete(
        f"/api/mandanten/{mandant_id}?bestaetigung=falsch").status_code == 422

    antwort = client.delete(f"/api/mandanten/{mandant_id}?bestaetigung=wegfall")
    assert antwort.status_code == 200, antwort.text
    assert client.get("/api/mandanten").json() == []
    assert client.get(f"/api/mandanten/{mandant_id}/plan").status_code == 404
    # Kurzname ist wieder frei
    neu = client.post("/api/mandanten", json={"name": "Neu", "kurzname": "wegfall"})
    assert neu.status_code == 201

    client.post("/api/benutzer", json={
        "email": "ml@example.com", "name": "M", "passwort": "geheim123",
        "rolle": "BEARBEITER", "mandanten_ids": [neu.json()["id"]]})
    client.get("/logout")
    client.post("/login", data={"email": "ml@example.com", "passwort": "geheim123"})
    assert client.delete(
        f"/api/mandanten/{neu.json()['id']}?bestaetigung=wegfall").status_code == 403


def test_dauer_import(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Dauer Import GmbH", "kurzname": "dauer-import"}
    ).json()["id"]
    csv_daten = (
        "Name;Art;Betrag;Intervall;Stichtag;GueltigVon;GueltigBis;Konto\n"
        "Miete Halle;Ausgabe;8900,00;monatlich;1;01.05.2026;;4210\n"
        "Leasing Stapler;Ausgabe;450,00;monatlich;15;01.05.2026;31.12.2027;9999\n"
        "Wartungspauschale;Einnahme;1250,00;vierteljährlich;10;01.05.2026;;8400\n"
    ).encode("utf-8")
    antwort = client.post(
        f"/api/mandanten/{mandant_id}/dauerbuchungen/import",
        files={"datei": ("dauer.csv", csv_daten, "text/csv")},
    )
    assert antwort.status_code == 200, antwort.text
    erg = antwort.json()
    assert erg["angelegt"] == 3
    assert erg["uebersprungen"] == 0
    assert any("9999" in w for w in erg["warnungen"])

    dauer = {d["name"]: d for d in
             client.get(f"/api/mandanten/{mandant_id}/dauerbuchungen").json()}
    assert dauer["Miete Halle"]["art"] == "KREDITOR"
    assert dauer["Miete Halle"]["konto_id"] is not None
    assert dauer["Miete Halle"]["gueltig_von"] == "2026-05-01"
    assert dauer["Leasing Stapler"]["konto_id"] is None         # 9999 unbekannt
    assert dauer["Leasing Stapler"]["gueltig_bis"] == "2027-12-31"
    assert dauer["Wartungspauschale"]["art"] == "DEBITOR"
    assert dauer["Wartungspauschale"]["intervall"] == "QUARTAL"

    # erneuter Import: alles Duplikate
    erg = client.post(
        f"/api/mandanten/{mandant_id}/dauerbuchungen/import",
        files={"datei": ("dauer.csv", csv_daten, "text/csv")},
    ).json()
    assert erg["angelegt"] == 0
    assert erg["uebersprungen"] == 3


def test_op_import_sammelspalte(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "OP Sammel GmbH", "kurzname": "op-sammel"}
    ).json()["id"]
    csv_daten = (
        "Art;Partner;Belegnummer;Rechnungsdatum;Faellig;Betrag;Konto;Notiz;Sammel\n"
        "FO;Debitoren-Sammeleingang KW 46;;;13.11.2026;10000,00;8400;je KW;ja\n"
        "FO;Kunde Einzeln;RE-2026-77;01.11.2026;20.11.2026;2500,00;8400;;\n"
    ).encode("utf-8")
    erg = client.post(
        f"/api/mandanten/{mandant_id}/posten/import",
        files={"datei": ("op.csv", csv_daten, "text/csv")},
    ).json()
    assert erg["angelegt"] == 2
    posten = {p["partner"]: p for p in client.get(f"/api/mandanten/{mandant_id}/posten").json()}
    assert posten["Debitoren-Sammeleingang KW 46"]["sammel"] is True
    assert posten["Kunde Einzeln"]["sammel"] is False

    # Überzahlung auf dem importierten Sammelposten funktioniert direkt
    pid = posten["Debitoren-Sammeleingang KW 46"]["id"]
    erg = client.post(f"/api/posten/{pid}/zahlungen", json={"betrag": 10500}).json()
    assert erg["status"] == "BEZAHLT"
    assert erg["unverteilt"] == 500.0   # kein weiterer Sammelposten vorhanden


def test_wizard_seite_nur_admin(client):
    _login(client)
    antwort = client.get("/mandanten-wizard")
    assert antwort.status_code == 200
    assert "Wizard" in antwort.text

    client.post("/api/benutzer", json={
        "email": "wz@example.com", "name": "W", "passwort": "geheim123",
        "rolle": "BEARBEITER", "mandanten_ids": []})
    client.get("/logout")
    client.post("/login", data={"email": "wz@example.com", "passwort": "geheim123"})
    antwort = client.get("/mandanten-wizard", follow_redirects=False)
    assert antwort.status_code == 303
    assert antwort.headers["location"] == "/"


def test_zahlung_mit_ist_buchung(client):
    _login(client)
    mandant_id = client.post(
        "/api/mandanten", json={"name": "Ist GmbH", "kurzname": "ist-gmbh"}
    ).json()["id"]
    konten = {k["nummer"]: k for k in
              client.get(f"/api/mandanten/{mandant_id}/konten").json()["konten"]}
    pid = client.post(f"/api/mandanten/{mandant_id}/posten", json={
        "art": "KREDITOR", "partner": "Stromwerk", "faellig_am": "2026-11-20",
        "betrag_brutto": 1000, "konto_id": konten["4240"]["id"]}).json()["id"]

    erg = client.post(f"/api/posten/{pid}/zahlungen", json={
        "betrag": 400, "datum": "2026-11-18",
        "bank_konto_id": konten["1200"]["id"]}).json()
    assert erg["status"] == "OFFEN"
    assert len(erg["buchungen"]) == 1
    assert erg["hinweise"] == []

    # Ist-Buchung in DATEV-Logik: Sachkonto an Bank am Zahltag
    from app import database
    db = database.SessionLocal()
    try:
        from app import models as m
        buchung = db.get(m.Buchung, erg["buchungen"][0])
        assert (buchung.konto_nr, buchung.gegenkonto_nr) == ("4240", "1200")
        assert float(buchung.betrag) == 400.0
        assert buchung.datum.isoformat() == "2026-11-18"
        assert buchung.posten_id == pid
    finally:
        db.close()

    # Posten ohne Konto: Zahlung wird angerechnet, aber mit Hinweis ohne Ist-Buchung
    ohne = client.post(f"/api/mandanten/{mandant_id}/posten", json={
        "art": "DEBITOR", "partner": "Kunde ohne Konto", "faellig_am": "2026-11-25",
        "betrag_brutto": 500}).json()["id"]
    erg = client.post(f"/api/posten/{ohne}/zahlungen", json={
        "betrag": 500, "bank_konto_id": konten["1200"]["id"]}).json()
    assert erg["status"] == "BEZAHLT"
    assert erg["buchungen"] == []
    assert any("kein Konto" in h for h in erg["hinweise"])
