"""Excel- und PDF-Export der 13-Wochen-Planung und des Soll-/Ist-Vergleichs.

Beide Exporte arbeiten auf dem Ergebnis von `berechne_plan` bzw.
`soll_ist_vergleich`, damit Anzeige und Bericht identisch rechnen.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

from .. import models

VERFAHREN_NAMEN = {
    "REGELMANDAT": "Regelmandat",
    "VORLAEUFIG": "vorläufiges Verfahren",
    "EROEFFNET": "eröffnetes Verfahren",
    "EIGENVERWALTUNG": "Eigenverwaltung",
}


def _auto(zeile: dict, tag: str, heute: str) -> float:
    """Anzeige-Logik: Ist bis heute, Plan ab morgen."""
    if tag <= heute:
        return zeile.get("ist", {}).get(tag, 0.0)
    return zeile.get("plan", {}).get(tag, 0.0)


def _agg(kinder: list[dict]) -> dict:
    agg: dict = {"plan": {}, "ist": {}, "basis": {}, "vortrag": {}}
    for kind in kinder:
        for feld in ("plan", "ist", "basis", "vortrag"):
            for tag, wert in kind.get(feld, {}).items():
                agg[feld][tag] = agg[feld].get(tag, 0.0) + wert
    return agg


def _wochenwert(zeile: dict, tage: list[str], heute: str, metrik: str) -> float | None:
    summe, belegt = 0.0, False
    for tag in tage:
        if metrik == "auto":
            wert = _auto(zeile, tag, heute)
        elif metrik == "plan":
            wert = zeile.get("plan", {}).get(tag, 0.0)
        elif metrik == "ist":
            if tag > heute:
                continue
            wert = zeile.get("ist", {}).get(tag, 0.0)
        else:  # delta
            if tag > heute:
                continue
            wert = zeile.get("ist", {}).get(tag, 0.0) - zeile.get("basis", {}).get(tag, 0.0)
        if wert:
            summe += wert
            belegt = True
    return summe if belegt else None


def _zellwert(zeile: dict, tage: list[str], heute: str, metrik: str) -> tuple[float | None, bool, bool]:
    """(Wert, enthält Vortrag, Vortrag trifft auf reguläre Planbuchung) einer
    Zelle über die Tage `tage`. „vortrag“ = Ist bis heute, Plan ab morgen
    zuzüglich vorgetragener, nicht gebuchter Dauerschuld-Zahlungen."""
    if metrik != "vortrag":
        return _wochenwert(zeile, tage, heute, metrik), False, False
    summe, belegt, mit_vortrag, gemischt = 0.0, False, False, False
    for tag in tage:
        if tag <= heute:
            wert = zeile.get("ist", {}).get(tag, 0.0)
        else:
            plan = zeile.get("plan", {}).get(tag, 0.0)
            vortrag = zeile.get("vortrag", {}).get(tag, 0.0)
            wert = plan + vortrag
            if vortrag:
                mit_vortrag = True
                gemischt = gemischt or bool(plan)
        if wert:
            summe += wert
            belegt = True
    return (summe if belegt else None), mit_vortrag, gemischt


_MONATSNAMEN = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August",
                "September", "Oktober", "November", "Dezember"]
_WOCHENTAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


def _monate(plan: dict) -> list[dict]:
    """Kalendermonate des Fensters (Rand-Monate nur mit den Fenstertagen)."""
    monate: list[dict] = []
    for tag in plan["tage"]:
        if not monate or monate[-1]["schluessel"] != tag[:7]:
            monate.append({
                "schluessel": tag[:7],
                "label": f"{_MONATSNAMEN[int(tag[5:7]) - 1]} {tag[:4]}",
                "tage": [],
            })
        monate[-1]["tage"].append(tag)
    return monate


def _spalten(plan: dict, ansicht: str) -> list[dict]:
    """Spalten der gewählten Ansicht: Kopfbeschriftung + zugehörige Tage."""
    if ansicht == "tage":
        return [
            {"kopf": f"{_WOCHENTAGE[date.fromisoformat(t).weekday()]}\n{t[8:10]}.{t[5:7]}.",
             "tage": [t]}
            for t in plan["tage"]
        ]
    if ansicht == "monate":
        return [
            {"kopf": f"{m['label']}\n{m['tage'][0][8:10]}.–{m['tage'][-1][8:10]}.",
             "tage": m["tage"]}
            for m in _monate(plan)
        ]
    return [
        {"kopf": f"{w['label']}\n{w['von'][8:10]}.{w['von'][5:7]}.–{w['bis'][8:10]}.{w['bis'][5:7]}.",
         "tage": w["tage"]}
        for w in plan["wochen"]
    ]


def _bestandsreihen(plan: dict, metrik: str) -> tuple[dict, dict]:
    """Liquidität/verfügbar – im Vortragsmodus inkl. vorgetragener Zahlungen."""
    b = plan["bestaende"]
    if metrik == "vortrag" and "liquiditaet_vortrag" in b:
        return b["liquiditaet_vortrag"], b["verfuegbar_vortrag"]
    return b["liquiditaet"], b["verfuegbar"]


_METRIK_TEXT = {
    "auto": "Ist bis heute, Plan ab morgen",
    "vortrag": "Ist bis heute, Plan ab morgen mit Vortrag nicht gebuchter Dauerschuld-Zahlungen",
    "plan": "Planwerte",
    "ist": "Ist-Werte bis heute",
    "delta": "Abweichung Ist − Soll (Soll-Referenz)",
}


def _soll_basis_text(plan: dict) -> str:
    if plan["vergleichsbasis"] != "SNAPSHOT":
        return "aktueller Plan (kein Snapshot)"
    snap = plan["snapshot"]
    art = "gewählte Soll-Referenz" if snap.get("referenz") else "jüngster eingefrorener Plan"
    return f"{art} vom {_datum_de(snap['stichtag'])}"


def _letzter_wert(werte: dict, tage: list[str]) -> float | None:
    """Letzter bekannter Wert innerhalb der Woche (Ist-Bestände enden heute)."""
    wert = None
    for tag in tage:
        w = werte.get(tag)
        if w is not None:
            wert = w
    return wert


def _datum_de(iso: str) -> str:
    return f"{iso[8:10]}.{iso[5:7]}.{iso[0:4]}"


def _de(wert: float | None) -> str:
    if wert is None or abs(wert) < 0.005:
        return ""
    return f"{wert:,.2f}".replace(",", " ").replace(".", ",").replace(" ", ".")


def _szenario_text(plan: dict) -> str | None:
    s = plan.get("szenario")
    if not s:
        return None
    teile = [f"Einzahlungen {s['ein_faktor']:.0f} %", f"variable Auszahlungen {s['aus_faktor']:.0f} %"]
    if s["debitoren_verzoegerung_tage"]:
        teile.append(f"Debitoren +{s['debitoren_verzoegerung_tage']} Tage")
    if s.get("regeln_anzahl"):
        teile.append(f"{s['regeln_anzahl']} Detailregel(n)")
    return f"{s['name']} ({', '.join(teile)})"


def _igeld_text(plan: dict) -> str | None:
    ig = plan.get("insolvenzgeld")
    if not ig:
        return None
    return (
        f"Zeitraum {_datum_de(ig['von'])} – {_datum_de(ig['bis'])}, "
        f"Entlastung im Fenster {_de(ig['entlastung_fenster'])} €"
    )


# --------------------------------------------------------------------- Excel

_FETT = Font(bold=True)
_GRUPPE_FUELLUNG = PatternFill("solid", fgColor="EDF1F7")
_SUMME_FUELLUNG = PatternFill("solid", fgColor="DCE6F2")
_BESTAND_FUELLUNG = PatternFill("solid", fgColor="F6F2E8")
_KOPF_FUELLUNG = PatternFill("solid", fgColor="C9D8EC")
_ZAHL_FORMAT = "#,##0.00;[Red]-#,##0.00"
_VORTRAG_FUELLUNG = PatternFill("solid", fgColor="FBE3B9")
_VORTRAG_GEMISCHT_FUELLUNG = PatternFill("darkUp", fgColor="E0A040", bgColor="FBE3B9")
_DUENN = Side(style="thin", color="D6DDE6")
_RAHMEN = Border(left=_DUENN, right=_DUENN, top=_DUENN, bottom=_DUENN)


def _zeilen_struktur(plan: dict) -> list[tuple[str, str, dict]]:
    """Liefert (typ, name, zeilen-dict) in Anzeige-Reihenfolge."""
    struktur: list[tuple[str, str, dict]] = []
    for gruppe in plan["zeilen"]:
        struktur.append(("gruppe", gruppe["name"], _agg(gruppe["kinder"])))
        for kind in gruppe["kinder"]:
            name = f"{kind['nummer']} {kind['name']}".strip()
            struktur.append(("konto", name, kind))
    s = plan["summen"]
    struktur.append(("summe", "Einzahlungen gesamt", s["einzahlungen"]))
    struktur.append(("summe", "Auszahlungen gesamt", s["auszahlungen"]))
    struktur.append(("summe", "Netto-Cashflow", s["netto"]))
    return struktur


def _perioden_blatt(ws, perioden: list[tuple[str, list[str]]], struktur: list, heute: str,
                    liquiditaet: dict, liquiditaet_text: str) -> None:
    """Blatt mit Plan/Ist/Δ je Periode (Woche bzw. Monat)."""
    ws.cell(row=1, column=1, value="Konto / BWA-Zeile").font = _FETT
    spalte = 2
    for kopf_text, _ in perioden:
        ws.merge_cells(start_row=1, start_column=spalte, end_row=1, end_column=spalte + 2)
        kopf = ws.cell(row=1, column=spalte, value=kopf_text)
        kopf.fill = _KOPF_FUELLUNG
        kopf.font = _FETT
        kopf.alignment = Alignment(horizontal="center")
        for i, unter in enumerate(("Plan", "Ist", "Δ")):
            z = ws.cell(row=2, column=spalte + i, value=unter)
            z.fill = _KOPF_FUELLUNG
            z.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(spalte + i)].width = 11
        spalte += 3
    ws.column_dimensions["A"].width = 38

    zeile_nr = 3
    for typ, name, daten in struktur:
        zelle = ws.cell(row=zeile_nr, column=1, value=name)
        if typ in ("gruppe", "summe"):
            zelle.font = _FETT
            zelle.fill = _GRUPPE_FUELLUNG if typ == "gruppe" else _SUMME_FUELLUNG
        spalte = 2
        for _, tage in perioden:
            for i, metrik in enumerate(("plan", "ist", "delta")):
                z = ws.cell(row=zeile_nr, column=spalte + i,
                            value=_wochenwert(daten, tage, heute, metrik))
                z.number_format = _ZAHL_FORMAT
                z.border = _RAHMEN
                if typ == "gruppe":
                    z.fill = _GRUPPE_FUELLUNG
                if typ == "summe":
                    z.fill = _SUMME_FUELLUNG
            spalte += 3
        zeile_nr += 1
    zelle = ws.cell(row=zeile_nr, column=1, value=liquiditaet_text)
    zelle.font = _FETT
    zelle.fill = _BESTAND_FUELLUNG
    spalte = 2
    for _, tage in perioden:
        z = ws.cell(row=zeile_nr, column=spalte, value=liquiditaet.get(tage[-1]))
        z.number_format = _ZAHL_FORMAT
        z.fill = _BESTAND_FUELLUNG
        spalte += 3
    ws.freeze_panes = "B3"


def plan_xlsx(plan: dict, sollist: dict | None, mandant: models.Mandant,
              metrik: str = "auto") -> bytes:
    wb = Workbook()
    heute = plan["heute"]
    # Blatt „Tage“ folgt der Werteauswahl, sofern sie eine Tagessicht ist
    tag_metrik = metrik if metrik in ("auto", "vortrag") else "auto"

    # ---- Info ----
    info = wb.active
    info.title = "Info"
    zeilen = [
        ("13-Wochen-Liquiditätsplanung", ""),
        ("Mandant", mandant.name),
        ("Aktenzeichen", mandant.aktenzeichen or "–"),
        ("Verfahrensstatus", VERFAHREN_NAMEN.get(mandant.verfahrensstatus, mandant.verfahrensstatus)),
        ("Insolvenz-Stichtag", _datum_de(mandant.insolvenz_stichtag.isoformat()) if mandant.insolvenz_stichtag else "–"),
        ("Zeitraum", f"{_datum_de(plan['start'])} – {_datum_de(plan['ende'])}"),
        ("Stand (heute)", _datum_de(heute)),
        ("Soll-Basis", _soll_basis_text(plan)),
        ("Szenario", _szenario_text(plan) or "Basisplan"),
        ("Insolvenzgeld", _igeld_text(plan) or "nicht aktiv"),
        ("Gesperrte Insolvenzforderungen (§ 38 InsO)", plan["gesperrte_insolvenzforderungen"]),
        ("Werte", f"{_METRIK_TEXT[tag_metrik]} (Blatt 'Tage'); Plan/Ist/Δ je Woche bzw. "
                  "Monat (Blätter 'Wochen', 'Monate'; Δ gegen die Soll-Basis)"),
    ]
    for i, (a, b) in enumerate(zeilen, start=1):
        info.cell(row=i, column=1, value=a).font = _FETT
        zelle = info.cell(row=i, column=2, value=b)
        if isinstance(b, (int, float)):
            zelle.number_format = _ZAHL_FORMAT
    info.column_dimensions["A"].width = 42
    info.column_dimensions["B"].width = 44

    struktur = _zeilen_struktur(plan)

    # ---- Tage ----
    ws = wb.create_sheet("Tage")
    ws.cell(row=1, column=1, value="Konto / BWA-Zeile").font = _FETT
    spalte = 2
    for woche in plan["wochen"]:
        von = spalte
        for tag in woche["tage"]:
            zelle = ws.cell(row=2, column=spalte, value=f"{tag[8:10]}.{tag[5:7]}.")
            zelle.fill = _KOPF_FUELLUNG
            zelle.alignment = Alignment(horizontal="center")
            if tag == heute:
                zelle.font = _FETT
            ws.column_dimensions[get_column_letter(spalte)].width = 10
            spalte += 1
        ws.merge_cells(start_row=1, start_column=von, end_row=1, end_column=spalte - 1)
        kopf = ws.cell(row=1, column=von, value=woche["label"])
        kopf.fill = _KOPF_FUELLUNG
        kopf.font = _FETT
        kopf.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 38

    zeile_nr = 3
    for typ, name, daten in struktur:
        zelle = ws.cell(row=zeile_nr, column=1, value=name)
        spalte = 2
        for woche in plan["wochen"]:
            for tag in woche["tage"]:
                wert, mit_vortrag, gemischt = _zellwert(daten, [tag], heute, tag_metrik)
                z = ws.cell(row=zeile_nr, column=spalte, value=wert if wert else None)
                z.number_format = _ZAHL_FORMAT
                z.border = _RAHMEN
                if typ == "gruppe":
                    z.fill = _GRUPPE_FUELLUNG
                if typ == "summe":
                    z.fill = _SUMME_FUELLUNG
                if mit_vortrag:
                    z.fill = _VORTRAG_GEMISCHT_FUELLUNG if (gemischt and typ == "konto") \
                        else _VORTRAG_FUELLUNG
                spalte += 1
        if typ in ("gruppe", "summe"):
            zelle.font = _FETT
            zelle.fill = _GRUPPE_FUELLUNG if typ == "gruppe" else _SUMME_FUELLUNG
        zeile_nr += 1

    b = plan["bestaende"]
    bestand_zeilen: list[tuple[str, dict[str, float | None]]] = [
        (f"{fk['nummer']} {fk['name']}", fk["bestand"]) for fk in b["finanzkonten"]
    ]
    liquiditaet, verfuegbar = _bestandsreihen(plan, tag_metrik)
    bestand_zeilen.append(("Liquidität (Bank + Kasse)", liquiditaet))
    if b["kreditlinien"]:
        bestand_zeilen.append(("Verfügbar (inkl. Kreditlinien)", verfuegbar))
    bestand_zeilen.append(("Warenbestand (nachrichtlich)", b["waren"]))
    for name, werte in bestand_zeilen:
        zelle = ws.cell(row=zeile_nr, column=1, value=name)
        zelle.fill = _BESTAND_FUELLUNG
        zelle.font = _FETT
        spalte = 2
        for woche in plan["wochen"]:
            for tag in woche["tage"]:
                wert = werte.get(tag)
                z = ws.cell(row=zeile_nr, column=spalte, value=wert)
                z.number_format = _ZAHL_FORMAT
                z.fill = _BESTAND_FUELLUNG
                z.border = _RAHMEN
                spalte += 1
        zeile_nr += 1
    ws.freeze_panes = "B3"

    # ---- Wochen und Monate (Plan/Ist/Δ gegen die Soll-Basis) ----
    wochen = [
        (f"{w['label']} ({w['von'][8:10]}.{w['von'][5:7]}.–{w['bis'][8:10]}.{w['bis'][5:7]}.)", w["tage"])
        for w in plan["wochen"]
    ]
    _perioden_blatt(wb.create_sheet("Wochen"), wochen, struktur, heute, liquiditaet,
                    "Liquidität zum Wochenende")
    monate = [(m["label"], m["tage"]) for m in _monate(plan)]
    _perioden_blatt(wb.create_sheet("Monate"), monate, struktur, heute, liquiditaet,
                    "Liquidität zum Monatsende")

    # ---- Soll-Ist (falls Snapshot vorhanden) ----
    if sollist and sollist.get("snapshot"):
        ws2 = wb.create_sheet("Soll-Ist")
        ws2.cell(row=1, column=1,
                 value=f"Soll-/Ist-Vergleich – "
                       f"{'gewählte Soll-Referenz' if sollist['snapshot'].get('referenz') else 'eingefrorener Plan'}"
                       f" vom {_datum_de(sollist['snapshot']['stichtag'])}"
                 ).font = _FETT
        ws2.cell(row=2, column=1, value="Zeile").font = _FETT
        spalte = 2
        for woche in sollist["wochen"]:
            ws2.merge_cells(start_row=2, start_column=spalte, end_row=2, end_column=spalte + 2)
            kopf = ws2.cell(row=2, column=spalte, value=woche["label"])
            kopf.fill = _KOPF_FUELLUNG
            kopf.font = _FETT
            kopf.alignment = Alignment(horizontal="center")
            for i, unter in enumerate(("Soll", "Ist", "Δ")):
                z = ws2.cell(row=3, column=spalte + i, value=unter)
                z.fill = _KOPF_FUELLUNG
                ws2.column_dimensions[get_column_letter(spalte + i)].width = 11
            spalte += 3
        ws2.column_dimensions["A"].width = 38
        zeile_nr = 4
        heute_iso = sollist["heute"]
        for z_ in sollist["zeilen"]:
            ws2.cell(row=zeile_nr, column=1, value=z_["name"])
            spalte = 2
            for woche in sollist["wochen"]:
                schluessel = f"{woche['jahr']}-{woche['kw']:02d}"
                soll = z_["plan"].get(schluessel)
                vergangen = woche["von"] <= heute_iso
                ist = z_["ist"].get(schluessel) if vergangen else None
                delta = None
                if vergangen and (soll is not None or ist is not None):
                    delta = (ist or 0.0) - (soll or 0.0)
                for i, wert in enumerate((soll, ist, delta)):
                    zc = ws2.cell(row=zeile_nr, column=spalte + i, value=wert)
                    zc.number_format = _ZAHL_FORMAT
                    zc.border = _RAHMEN
                spalte += 3
            zeile_nr += 1
        ws2.freeze_panes = "B4"

    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


# ----------------------------------------------------------------------- PDF

# Spalten je Blatt (Querformat A4); breitere Pläne laufen nach rechts auf
# Folgeblätter weiter
SPALTEN_JE_BLATT = {"tage": 14, "wochen": 13, "monate": 13}
_VORTRAG_FARBE = colors.HexColor("#FBE3B9")
_VORTRAG_GEMISCHT_FARBE = colors.HexColor("#F2C27A")


def plan_pdf(plan: dict, mandant: models.Mandant, ansicht: str = "wochen",
             metrik: str = "auto") -> bytes:
    heute = plan["heute"]
    puffer = BytesIO()
    titel = f"{len(plan['wochen'])}-Wochen-Liquiditätsplanung"
    dokument = SimpleDocTemplate(
        puffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=f"{titel} {mandant.name}",
    )
    titel_stil = ParagraphStyle("titel", fontName="Helvetica-Bold", fontSize=13, leading=17,
                                spaceAfter=2)
    meta_stil = ParagraphStyle("meta", fontName="Helvetica", fontSize=8, textColor=colors.HexColor("#444444"))
    blatt_stil = ParagraphStyle("blatt", fontName="Helvetica-Bold", fontSize=9, spaceAfter=3)

    elemente = [Paragraph(titel, titel_stil)]
    meta = (
        f"{mandant.name}"
        + (f" · {mandant.aktenzeichen}" if mandant.aktenzeichen else "")
        + f" · {VERFAHREN_NAMEN.get(mandant.verfahrensstatus, mandant.verfahrensstatus)}"
        + f" · Zeitraum {_datum_de(plan['start'])} – {_datum_de(plan['ende'])}"
        + f" · Stand {_datum_de(heute)}"
        + f" · Soll-Basis: {_soll_basis_text(plan)}"
    )
    elemente.append(Paragraph(meta, meta_stil))
    elemente.append(Paragraph(f"Werte: {_METRIK_TEXT[metrik]}", meta_stil))
    szenario_text = _szenario_text(plan)
    if szenario_text:
        elemente.append(Paragraph(f"<b>Szenario:</b> {szenario_text}", meta_stil))
    igeld_text = _igeld_text(plan)
    if igeld_text:
        elemente.append(Paragraph(f"Insolvenzgeld: {igeld_text}", meta_stil))
    if plan["gesperrte_insolvenzforderungen"]:
        elemente.append(Paragraph(
            f"Zahlungsgesperrte Insolvenzforderungen (§ 38 InsO), nicht im Plan: "
            f"{_de(plan['gesperrte_insolvenzforderungen'])} €", meta_stil))
    if metrik == "vortrag" and plan.get("vortraege"):
        teile = [
            f"{v['name']} (fällig {_datum_de(v['faellig'])}, offen {_de(abs(v['betrag']))} €"
            + (f" von {_de(abs(v['soll']))} €" if v["gezahlt"] else "") + ")"
            for v in plan["vortraege"]
        ]
        elemente.append(Paragraph(
            f"Vorgetragen auf {_datum_de(plan['vortraege'][0]['ziel'])}: " + "; ".join(teile),
            meta_stil))
    elemente.append(Spacer(0, 4 * mm))

    spalten = _spalten(plan, ansicht)

    # Zeilen über alle Spalten bestimmen, damit jedes Blatt dieselben Zeilen zeigt
    zeilen: list[tuple[str, str, list]] = []   # (typ, name, [(wert, vortrag, gemischt)])
    for typ, name, zeilen_daten in _zeilen_struktur(plan):
        zellen = [_zellwert(zeilen_daten, sp["tage"], heute, metrik) for sp in spalten]
        if typ == "konto" and not any(z[0] is not None for z in zellen):
            continue
        zeilen.append((typ, name, zellen))
    b = plan["bestaende"]
    liquiditaet, verfuegbar = _bestandsreihen(plan, metrik)
    for fk in b["finanzkonten"]:
        zeilen.append(("bestand", f"{fk['nummer']} {fk['name']}",
                       [(_letzter_wert(fk["bestand"], sp["tage"]), False, False) for sp in spalten]))
    zeilen.append(("bestand", "Liquidität (Bank + Kasse)",
                   [(liquiditaet.get(sp["tage"][-1]), False, False) for sp in spalten]))
    if b["kreditlinien"]:
        zeilen.append(("bestand", "Verfügbar (inkl. Kreditlinien)",
                       [(verfuegbar.get(sp["tage"][-1]), False, False) for sp in spalten]))
    zeilen.append(("bestand", "Warenbestand (nachrichtlich)",
                   [(b["waren"].get(sp["tage"][-1]), False, False) for sp in spalten]))

    je_blatt = SPALTEN_JE_BLATT[ansicht]
    bloecke = [range(i, min(i + je_blatt, len(spalten))) for i in range(0, len(spalten), je_blatt)]
    erste = 62 * mm
    spalten_breite = (dokument.width - erste) / min(je_blatt, len(spalten))

    for blatt_nr, block in enumerate(bloecke, start=1):
        if blatt_nr > 1:
            elemente.append(PageBreak())
        if len(bloecke) > 1:
            von, bis = spalten[block[0]]["tage"][0], spalten[block[-1]]["tage"][-1]
            elemente.append(Paragraph(
                f"{titel} · {mandant.name} · Blatt {blatt_nr} von {len(bloecke)}: "
                f"{_datum_de(von)} – {_datum_de(bis)}", blatt_stil))
        daten: list[list[str]] = [["Konto / BWA-Zeile"] + [spalten[i]["kopf"] for i in block]]
        stile = [
            ("FONTSIZE", (0, 0), (-1, -1), 6.5),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C9D8EC")),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D6DDE6")),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 1.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ]
        for typ, name, zellen in zeilen:
            nr = len(daten)
            daten.append([("    " + name) if typ == "konto" else name]
                         + [_de(zellen[i][0]) for i in block])
            hintergrund = {"gruppe": "#EDF1F7", "summe": "#DCE6F2", "bestand": "#F6F2E8"}.get(typ)
            if hintergrund:
                stile.append(("BACKGROUND", (0, nr), (-1, nr), colors.HexColor(hintergrund)))
            if typ in ("gruppe", "summe"):
                stile.append(("FONTNAME", (0, nr), (-1, nr), "Helvetica-Bold"))
            elif typ == "bestand":
                stile.append(("FONTNAME", (0, nr), (0, nr), "Helvetica-Bold"))
            for spalte, i in enumerate(block, start=1):
                wert, mit_vortrag, gemischt = zellen[i]
                if mit_vortrag:
                    farbe = _VORTRAG_GEMISCHT_FARBE if (gemischt and typ == "konto") else _VORTRAG_FARBE
                    stile.append(("BACKGROUND", (spalte, nr), (spalte, nr), farbe))
                if wert is not None and wert < -0.005:
                    stile.append(("TEXTCOLOR", (spalte, nr), (spalte, nr), colors.HexColor("#B3362C")))
        tabelle = Table(daten, colWidths=[erste] + [spalten_breite] * len(block), repeatRows=1)
        tabelle.setStyle(TableStyle(stile))
        elemente.append(tabelle)

    einheit = {"tage": ("Tageswerte", "Tagesende"), "wochen": ("Wochenwerte", "Wochenende"),
               "monate": ("Monatswerte", "Monatsende")}[ansicht]
    hinweis = f"{einheit[0]}: {_METRIK_TEXT[metrik]}. Bestände jeweils zum {einheit[1]}."
    if metrik == "vortrag":
        hinweis += (" Hell hinterlegt: vorgetragene Zahlung; dunkler: Vortrag trifft auf eine "
                    "reguläre Planbuchung desselben Kontos (Werte addiert).")
    elemente.append(Spacer(0, 3 * mm))
    elemente.append(Paragraph(hinweis + " Erstellt mit LiquiPlanung.", meta_stil))
    dokument.build(elemente)
    return puffer.getvalue()
