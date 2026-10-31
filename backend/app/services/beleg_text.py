"""Text und Seitenbilder aus Belegen (Bestelldokumente) für das KI-Auslesen.

Unterstützt PDF (Textebene, sonst Texterkennung), Bilder (PNG/JPEG/TIFF/WebP,
Texterkennung), DOCX, XLSX, E-Mails (.eml inkl. Anhängen) und Textdateien.
Bewusst ohne Kontoauszugs-Heuristik: es werden nur die generischen Bausteine
aus `pdf_kontoauszug` verwendet (pypdf-Leser, tesseract, pdfium-Sperre).

Fehler führen nie zu einer Ausnahme, sondern zu Warnungen im Ergebnis.
"""
from __future__ import annotations

import base64
import email
import email.policy
import functools
import html
import logging
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from pathlib import PurePosixPath
from xml.etree import ElementTree
from xml.parsers import expat

from . import pdf_kontoauszug as _pk

log = logging.getLogger(__name__)

BELEG_MAX_SEITEN = 5          # Seiten für Text/OCR
VISION_MAX_SEITEN = 3         # Seitenbilder für Vision-Modelle
VISION_DPI = 150
VISION_MAX_KANTE = 2000       # px (OpenAI-kompatible Dienste skalieren größer ohnehin herunter)
OCR_MIN_KANTE = 1600          # kleinere Bilder vor der Texterkennung vergrößern
OCR_MAX_KANTE = 4000
TEXT_MAX = 200_000            # Schutz vor riesigen Tabellen/Mails (KI kürzt weiter)
BILD_MAX_PIXEL = 40_000_000   # je Einzelbild nach dem Dekodieren (A4 mit 600 dpi ≈ 35 MP)
BILD_MIN_KANTE = 300          # kleinere Bilder in Mails (Logos, Symbole) auslassen
BILD_GROSS_FLAECHE = 250_000  # eingebettete Mailbilder neben echten Anhängen (Foto, Screenshot)
EML_MAX_ANHAENGE = 10         # ausgewertete Anhänge je E-Mail
# Office-Dateien sind ZIP-Archive: Grenzen gegen Archivbomben (entpackt)
XML_MAX_BYTES = 20 * 1024 * 1024      # je ausgewertetem Teil
ZIP_SUMME_MAX = 50 * 1024 * 1024      # alle ausgewerteten Teile zusammen (E-Mail: alle Office-Anhänge)
ZIP_MAX_EINTRAEGE = 2000
ZIP_MAX_ELEMENTE = 500_000            # XML-Elemente der ganz eingelesenen Teile zusammen, XLSX-Blätter je Blatt
XML_MAX_NAMEN = 5_000                 # verschiedene Element- und Attributnamen (mit Namensraum) je Teil
XML_MAX_ATTRIBUTE = 1_000             # Attribute je Element
XML_MAX_NS_URI = 256                  # Zeichen je Namensraum-URI (OOXML: unter 100)
DOCX_MAX_KOPF_FUSS = 10               # je Kopf- bzw. Fußzeilenteile
XLSX_MAX_BLAETTER = 10
XLSX_MAX_ZEILEN = 1000
XLSX_MAX_SPALTEN = 50

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EML = "message/rfc822"
TXT = "text/plain"
BILD_MIME = {"image/png", "image/jpeg", "image/tiff", "image/webp", "image/gif", "image/bmp"}

_ENDUNG_MIME = {
    ".pdf": PDF, ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".tif": "image/tiff", ".tiff": "image/tiff", ".webp": "image/webp",
    ".docx": DOCX, ".xlsx": XLSX, ".eml": EML, ".txt": TXT, ".csv": TXT,
}


@dataclass
class BelegText:
    text: str
    quelle: str                 # "TEXT" | "OCR" | "LEER"
    seiten: int
    warnungen: list[str] = field(default_factory=list)


# ------------------------------------------------------------------ Hilfen

def ist_bild(mime: str | None) -> bool:
    return _mime_normal(mime) in BILD_MIME


def _mime_normal(mime: str | None) -> str:
    """„Text/Plain; charset=utf-8“ → „text/plain“; image/jpg → image/jpeg."""
    m = (mime or "").split(";")[0].strip().lower()
    return {"image/jpg": "image/jpeg", "image/pjpeg": "image/jpeg",
            "image/tif": "image/tiff", "application/x-pdf": PDF}.get(m, m)


def mime_bestimmen(daten: bytes, mime: str | None = None, dateiname: str | None = None) -> str:
    """MIME-Typ für die Textgewinnung: Inhalt (Magic Bytes) vor Angabe vor Endung."""
    if b"%PDF-" in daten[:1024]:
        return PDF
    if daten.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if daten.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if daten[:4] in (b"II*\x00", b"MM\x00*"):
        return "image/tiff"
    if daten[:4] == b"RIFF" and daten[8:12] == b"WEBP":
        return "image/webp"
    m = _mime_normal(mime)
    if daten[:4] == b"PK\x03\x04" and m not in (DOCX, XLSX):
        try:
            with zipfile.ZipFile(BytesIO(daten)) as z:
                namen = z.namelist()
            if any(n.startswith("word/") for n in namen):
                return DOCX
            if any(n.startswith("xl/") for n in namen):
                return XLSX
        except (zipfile.BadZipFile, ValueError, OSError):
            pass
    if m and m != "application/octet-stream":
        return m
    endung = PurePosixPath(str(dateiname or "").replace("\\", "/")).suffix.lower()
    return _ENDUNG_MIME.get(endung, m or "application/octet-stream")


def _dekodiere(daten: bytes) -> str:
    for kodierung in ("utf-8-sig", "cp1252"):
        try:
            return daten.decode(kodierung)
        except UnicodeDecodeError:
            continue
    return daten.decode("latin-1", "replace")


def _aufraeumen(text: str) -> str:
    """Leerraum verdichten: Zeilenenden, lange Leerzeichenfolgen (Layout-Modus),
    mehrfache Leerzeilen. Spaltentrennung bleibt als drei Leerzeichen sichtbar."""
    zeilen = []
    leer = 0
    for zeile in text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "").split("\n"):
        zeile = re.sub(r"[ \t ]{4,}", "   ", zeile.rstrip())
        if not zeile.strip():
            leer += 1
            if leer > 1:
                continue
            zeile = ""
        else:
            leer = 0
        zeilen.append(zeile)
    return "\n".join(zeilen).strip("\n")


@functools.cache
def _sprache() -> str:
    return _pk._ocr_sprache()


def _tsv_zeilen(tsv: str) -> list[str]:
    """tesseract-TSV → Textzeilen (gruppiert nach Block/Absatz/Zeile, Wörter links→rechts)."""
    zeilen: dict[tuple, list[tuple[int, str]]] = {}
    for roh in tsv.splitlines()[1:]:
        t = roh.split("\t")
        if len(t) < 12 or t[0] != "5" or not t[11].strip():
            continue
        try:
            schluessel = (int(t[1]), int(t[2]), int(t[3]), int(t[4]))
            links = int(t[6])
        except ValueError:
            continue
        zeilen.setdefault(schluessel, []).append((links, t[11].strip()))
    return [" ".join(w for _, w in sorted(woerter)) for woerter in zeilen.values()]


def _ohne_ocr_warnung() -> str:
    return ("Das Dokument enthält keinen auslesbaren Text (Scan bzw. Foto). Die dafür "
            "nötige Texterkennung (tesseract) ist auf dem Server nicht installiert.")


# ------------------------------------------------------------------ PDF

def pdf_text(daten: bytes, max_seiten: int = BELEG_MAX_SEITEN) -> tuple[list[str], int]:
    """Text je Seite (pypdf, Layout-Modus mit Rückfall auf den normalen Modus)
    und Gesamtzahl der Seiten. Wirft bei defekten PDFs."""
    leser = _pk._pdf_leser(daten)
    if leser.is_encrypted:
        try:
            leser.decrypt("")  # nur Besitzerpasswort: mit leerem Benutzerpasswort lesbar
        except Exception as e:  # noqa: BLE001 – Meldung folgt beim Lesen
            log.debug("PDF-Entschlüsselung fehlgeschlagen: %s", e)
    seiten = leser.pages
    anzahl = len(seiten)
    texte = []
    for nr in range(min(anzahl, max_seiten)):
        seite = seiten[nr]
        text = ""
        try:
            text = seite.extract_text(extraction_mode="layout") or ""
        except Exception:  # noqa: BLE001 – Layout-Modus scheitert an manchen Fonts
            text = ""
        if not text.strip():
            try:
                text = seite.extract_text() or ""
            except Exception:  # noqa: BLE001
                text = ""
        texte.append(text)
    return texte, anzahl


def _pdf(daten: bytes, max_seiten: int) -> BelegText:
    warnungen: list[str] = []
    try:
        seiten, anzahl = pdf_text(daten, max_seiten)
    except Exception as e:  # noqa: BLE001 – defekte/verschlüsselte PDFs sauber melden
        seiten, anzahl = [], 0
        warnungen.append(f"PDF-Text nicht lesbar: {e}")
    if anzahl > max_seiten:
        warnungen.append(f"Nur die ersten {max_seiten} von {anzahl} Seiten ausgewertet.")
    if seiten and not _pk._ohne_textebene(seiten):
        text = "\n\n".join(_aufraeumen(s) for s in seiten if s.strip())
        return BelegText(text, "TEXT", anzahl, warnungen)

    rest = _aufraeumen("\n".join(seiten))  # evtl. wenige Zeichen (Stempel, Kopfzeile)
    if not _pk.ocr_verfuegbar():
        warnungen.append(_ohne_ocr_warnung())
        return BelegText(rest, "LEER", anzahl, warnungen)
    try:
        ocr_seiten = _pk._ocr_zeilen(daten, bis=max_seiten)
        if not anzahl:
            anzahl = len(ocr_seiten)
        text = "\n\n".join(
            "\n".join(_pk._zeilentext(z) for z in seite) for seite in ocr_seiten
        ).strip()
    except Exception as e:  # noqa: BLE001
        warnungen.append(f"Texterkennung fehlgeschlagen: {e}")
        return BelegText(rest, "LEER", anzahl, warnungen)
    if not text:
        warnungen.append("Per Texterkennung wurde kein Text erkannt.")
        return BelegText(rest, "LEER", anzahl, warnungen)
    warnungen.append("Text per Texterkennung (OCR) gelesen – Lesefehler möglich.")
    return BelegText(_aufraeumen(text), "OCR", anzahl, warnungen)


def pdf_seiten_bilder(daten: bytes, max_seiten: int = VISION_MAX_SEITEN,
                      dpi: int = VISION_DPI, graustufen: bool = False) -> list:
    """PDF-Seiten als PIL-Bilder (pypdfium2, unter der gemeinsamen pdfium-Sperre).
    Übergroße Seiten werden gleich mit höchstens VISION_MAX_KANTE Pixeln gerendert."""
    import pypdfium2 as pdfium

    with _pk._PDFIUM_SPERRE:
        dokument = pdfium.PdfDocument(daten)
        try:
            bilder = []
            for nr in range(min(max_seiten, len(dokument))):
                seite = dokument[nr]
                try:
                    skala = _pk._render_skala(seite, dpi, max_kante=VISION_MAX_KANTE)
                    bilder.append(seite.render(scale=skala, grayscale=graustufen).to_pil())
                finally:
                    seite.close()
            return bilder
        finally:
            dokument.close()


# ------------------------------------------------------------------ Bilder

def _zielgroesse(groesse: tuple[int, int], max_kante: int) -> tuple[int, int]:
    breite, hoehe = groesse
    faktor = min(1.0, max_kante / max(breite, hoehe, 1))
    return max(1, round(breite * faktor)), max(1, round(hoehe * faktor))


def _bild_frames(daten: bytes, max_seiten: int, max_kante: int = VISION_MAX_KANTE):
    """Einzelbilder (mehrseitige TIFFs: je Seite eines) nacheinander, auf höchstens
    `max_kante` Pixel verkleinert, ausgerichtet nach EXIF, als RGB bzw. Graustufen
    ohne Transparenz. Jeder Frame wird einzeln dekodiert und sofort verkleinert;
    Frames über BILD_MAX_PIXEL (Dekompressionsbomben, z. B. Fax-TIFFs) → ValueError."""
    from PIL import Image, ImageOps, ImageSequence

    bild = Image.open(BytesIO(daten))
    grenze = 1 if bild.format == "MPO" else max_seiten   # MPO: nur das Hauptbild
    for nr, frame in enumerate(ImageSequence.Iterator(bild)):
        if nr >= grenze:
            break
        if bild.format in ("JPEG", "MPO"):
            frame.draft(None, _zielgroesse(frame.size, max_kante))   # verkleinert dekodieren
        breite, hoehe = frame.size
        if breite * hoehe > BILD_MAX_PIXEL:
            raise ValueError(f"Bild zu groß ({breite} × {hoehe} Pixel)")
        einzel = frame
        if einzel.mode == "1":
            einzel = einzel.convert("L")
        if einzel.mode in ("RGBA", "LA", "P", "PA") or "transparency" in einzel.info:
            einzel = einzel.convert("RGBA")
        elif einzel.mode not in ("RGB", "L"):
            einzel = einzel.convert("RGB")
        if max(einzel.size) > max_kante:
            einzel = einzel.resize(_zielgroesse(einzel.size, max_kante),
                                   Image.Resampling.BICUBIC, reducing_gap=2.0)
        if einzel is frame:
            einzel = frame.copy()   # der Iterator verwendet ein Objekt für alle Frames
        ImageOps.exif_transpose(einzel, in_place=True)
        if einzel.mode == "RGBA":
            hintergrund = Image.new("RGB", einzel.size, "white")
            hintergrund.paste(einzel, mask=einzel.getchannel("A"))
            einzel = hintergrund
        yield einzel


def _bild_anzahl(daten: bytes) -> int:
    try:
        from PIL import Image

        bild = Image.open(BytesIO(daten))
        if bild.format == "MPO":   # Handyfotos: Folgebilder sind Tiefen-/HDR-Karten, keine Seiten
            return 1
        return max(1, getattr(bild, "n_frames", 1))
    except Exception:  # noqa: BLE001
        return 1


def _png(bild) -> bytes:
    puffer = BytesIO()
    bild.save(puffer, "PNG")
    return puffer.getvalue()


def _kleines_bild(daten: bytes) -> bool:
    """Längste Kante unter BILD_MIN_KANTE (Logo, Symbol) oder nicht lesbar."""
    try:
        from PIL import Image

        return max(Image.open(BytesIO(daten)).size) < BILD_MIN_KANTE
    except Exception:  # noqa: BLE001
        return True


def _grosses_bild(daten: bytes) -> bool:
    """Foto bzw. Screenshot statt Logo oder Signatur-Banner: kürzere Kante mindestens
    BILD_MIN_KANTE, Fläche mindestens BILD_GROSS_FLAECHE (Banner sind flach und breit)."""
    try:
        from PIL import Image

        breite, hoehe = Image.open(BytesIO(daten)).size
    except Exception:  # noqa: BLE001
        return False
    return min(breite, hoehe) >= BILD_MIN_KANTE and breite * hoehe >= BILD_GROSS_FLAECHE


def bild_normalisieren(bild, max_kante: int = VISION_MAX_KANTE):
    """PIL-Bild auf höchstens `max_kante` Pixel (längste Seite) verkleinern."""
    if max(bild.size) > max_kante:
        bild = bild.copy()
        bild.thumbnail((max_kante, max_kante))
    return bild


def _ocr_bild(bild) -> str:
    from PIL import Image

    grau = bild.convert("L")
    kante = max(grau.size) or 1
    if kante < OCR_MIN_KANTE:
        faktor = min(3.0, OCR_MIN_KANTE / kante)
        grau = grau.resize((round(grau.width * faktor), round(grau.height * faktor)),
                           Image.Resampling.LANCZOS)
    elif kante > OCR_MAX_KANTE:
        grau.thumbnail((OCR_MAX_KANTE, OCR_MAX_KANTE))
    return "\n".join(_tsv_zeilen(_pk._tesseract_tsv(_png(grau), _sprache())))


def _bild(daten: bytes, max_seiten: int) -> BelegText:
    warnungen: list[str] = []
    anzahl = _bild_anzahl(daten)
    if not _pk.ocr_verfuegbar():
        warnungen.append(_ohne_ocr_warnung())
        return BelegText("", "LEER", anzahl, warnungen)
    try:
        texte = [_ocr_bild(f) for f in _bild_frames(daten, max_seiten, OCR_MAX_KANTE)]
    except Exception as e:  # noqa: BLE001 – defekte Bilder, Dekompressionsbomben, Zeitüberschreitung
        warnungen.append(f"Bild nicht lesbar: {e}")
        return BelegText("", "LEER", anzahl, warnungen)
    if anzahl > max_seiten:
        warnungen.append(f"Nur die ersten {max_seiten} von {anzahl} Seiten ausgewertet.")
    text = _aufraeumen("\n\n".join(t for t in texte if t.strip()))
    if not text:
        warnungen.append("Per Texterkennung wurde kein Text erkannt.")
        return BelegText("", "LEER", anzahl, warnungen)
    warnungen.append("Text per Texterkennung (OCR) gelesen – Lesefehler möglich.")
    return BelegText(text, "OCR", anzahl, warnungen)


# ------------------------------------------------------------------ Office

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_SML_ZEILE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main}row"   # expat-Schreibweise
_NS_DEKLARATION = re.compile(rb"xmlns(?::[^\s=]*)?\s*=\s*([\"'])")
_ZUWEISUNG = re.compile(rb"=\s*[\"']")


def _mb(n: int) -> str:
    return f"{n // (1024 * 1024)} MB"


def _zahl(n: int) -> str:
    return f"{n:,}".replace(",", ".")


def _xml_roh_pruefen(name: str, roh: bytes) -> None:
    """Prüfungen auf dem Rohtext, vor jedem Parser: keine DTD (Entity-Expansion, über
    Standardattribute versteckte Namensräume), Namensraum-URIs höchstens XML_MAX_NS_URI
    Zeichen, höchstens XML_MAX_ATTRIBUTE Attribute je Element. Beides vervielfacht sonst
    schon in expat den Speicher: jeder Name wird samt URI aufgelöst („{uri}name“).
    Attributwerte enthalten kein „<“: alle Attribute eines Tags stehen vor dem nächsten „<“."""
    if roh[:2] in (b"\xff\xfe", b"\xfe\xff") or b"\x00" in roh[:2]:   # UTF-16, Erkennung wie expat
        be = roh[:2] == b"\xfe\xff" or roh[:1] == b"\x00"
        roh = roh.decode("utf-16-be" if be else "utf-16-le", "replace").encode("utf-8")
    if b"<!ENTITY" in roh:
        raise ValueError(f"{name}: Entity-Deklaration nicht zulässig")
    if b"<!DOCTYPE" in roh:
        raise ValueError(f"{name}: Dokumenttyp-Deklaration (DTD) nicht zulässig")
    for treffer in _NS_DEKLARATION.finditer(roh):
        ende = roh.find(treffer.group(1), treffer.end())
        if (ende if ende >= 0 else len(roh)) - treffer.end() > XML_MAX_NS_URI:
            raise ValueError(f"{name}: Namensraum-Angabe zu lang (höchstens {XML_MAX_NS_URI} Zeichen)")
    # nur lange Abschnitte von „<“ bis zum nächsten „<“ können so viele Attribute enthalten
    for treffer in re.finditer(rb"<[^<]{%d,}" % XML_MAX_ATTRIBUTE, roh):
        if len(_ZUWEISUNG.findall(roh, treffer.start(), treffer.end())) > XML_MAX_ATTRIBUTE:
            raise ValueError(f"{name}: zu viele Attribute in einem Element "
                             f"(höchstens {_zahl(XML_MAX_ATTRIBUTE)})")


def _xml_vorpruefen(name: str, roh: bytes, grenze: int, blatt: bool = False) -> int:
    """Vorprüfung mit expat ohne Baum: höchstens XML_MAX_NAMEN verschiedene aufgelöste
    Element- und Attributnamen (ElementTree legt jeden zweimal in einem Zwischenspeicher
    ab) und höchstens `grenze` Elemente, die der Parser gleichzeitig hält. Das sind alle,
    bei streamend gelesenen XLSX-Blättern (`blatt`) die außerhalb von Zeilen (samt den
    geleerten Zeilen) plus die größte Zeile: openpyxl leert jede Zeile nach dem Lesen,
    andere Elemente nicht. Liefert diese Zahl. Ein Syntaxfehler beendet die Prüfung
    still, dort bricht auch der eigentliche Parser ab."""
    namen: set[str] = set()
    elemente = zeile = groesste = tiefe = 0

    def zu_viele() -> None:
        raise ValueError(f"zu umfangreich zum Auslesen (über {_zahl(ZIP_MAX_ELEMENTE)} XML-Elemente)")

    def start(tag: str, attribute: list[str]) -> None:
        nonlocal elemente, zeile, tiefe
        namen.add(tag)
        if attribute:
            namen.update(attribute[::2])
        if len(namen) > XML_MAX_NAMEN:
            raise ValueError(f"{name}: zu viele verschiedene XML-Namen (über {_zahl(XML_MAX_NAMEN)})")
        if tiefe:
            zeile += 1
            tiefe += tag == _SML_ZEILE
            if elemente + zeile > grenze:
                zu_viele()
        else:
            elemente += 1
            if blatt and tag == _SML_ZEILE:
                tiefe = 1
            if elemente + groesste > grenze:
                zu_viele()

    def ende(tag: str) -> None:
        nonlocal zeile, groesste, tiefe
        if tiefe and tag == _SML_ZEILE:
            tiefe -= 1
            if not tiefe:
                groesste, zeile = max(groesste, zeile), 0

    def dtd(*_args) -> None:
        raise ValueError(f"{name}: Dokumenttyp-Deklaration (DTD) nicht zulässig")

    parser = expat.ParserCreate(namespace_separator="}")   # Trenner wie ElementTree
    parser.ordered_attributes = True
    parser.StartElementHandler = start
    parser.StartDoctypeDeclHandler = dtd
    if blatt:
        parser.EndElementHandler = ende
    try:
        parser.Parse(roh, True)
    except expat.ExpatError:
        pass
    return elemente + max(groesste, zeile)


def _zip_groessen(z: zipfile.ZipFile, infos: list[zipfile.ZipInfo]) -> None:
    if len(z.infolist()) > ZIP_MAX_EINTRAEGE:
        raise ValueError(f"zu viele Teile im Archiv ({len(z.infolist())}, höchstens {ZIP_MAX_EINTRAEGE})")
    if (any(i.file_size > XML_MAX_BYTES for i in infos)
            or sum(i.file_size for i in infos) > ZIP_SUMME_MAX):
        raise ValueError(f"Inhalt entpackt zu groß zum Auslesen (höchstens {_mb(ZIP_SUMME_MAX)}, "
                         f"je Teil {_mb(XML_MAX_BYTES)})")


def _teile_pruefen(z: zipfile.ZipFile, infos: list[zipfile.ZipInfo], elemente: int = 0,
                   blaetter: frozenset[str] = frozenset()) -> int:
    """Rohtext- und expat-Prüfung der Teile; liefert die Elementsumme der ganz eingelesenen
    Teile (`blaetter` zählen nur je Blatt)."""
    for info in infos:
        if info.is_dir():
            continue
        roh = z.read(info)   # liefert höchstens file_size Bytes
        _xml_roh_pruefen(info.filename, roh)
        if info.filename in blaetter:
            _xml_vorpruefen(info.filename, roh, ZIP_MAX_ELEMENTE, blatt=True)
        else:
            elemente += _xml_vorpruefen(info.filename, roh, ZIP_MAX_ELEMENTE - elemente)
    return elemente


def _zip_pruefen(z: zipfile.ZipFile, namen: list[str] | None = None) -> None:
    """Schutz vor Archiv- und XML-Bomben vor dem Parsen der Teile `namen` (sonst aller
    Einträge): Zahl der Einträge, entpackte Größe je Teil und gesamt, Rohtext (DTD,
    Namensräume, Attribute) und expat-Vorprüfung (Namen, Elemente zusammen)."""
    infos = z.infolist() if namen is None else [z.getinfo(n) for n in namen]
    _zip_groessen(z, infos)
    _teile_pruefen(z, infos)


_CT_BLATT = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
_CT_MAPPE = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml",
    "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    "application/vnd.ms-excel.template.macroEnabled.main+xml",
}
# liest openpyxl (read_only) immer ganz bzw. als Ganzes, wenn vorhanden
_XLSX_FEST = {"[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml", "docProps/core.xml",
              "docProps/custom.xml", "xl/theme/theme1.xml"}


def _paketteil(name: str) -> bool:
    return name == "[Content_Types].xml" or name.endswith(".rels")


def _lokal(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _xlsx_blaetter(z: zipfile.ZipFile) -> frozenset[str]:
    """Arbeitsblätter, die openpyxl im read_only-Modus nur streamend liest: Ziele von
    Blatt-Beziehungen einer Mappe, die nirgends sonst vorkommen. Diagrammblätter,
    Zeichnungen, Diagramme, geteilte Zeichenketten usw. liest openpyxl ganz ein, auch
    unter einem Blatt-Pfad. Die Datei bestimmt, welcher Teil was ist – daher wie openpyxl
    auflösen (Elemente nur nach lokalem Namen, Ziele relativ zum Elternordner) und im
    Zweifel nichts ausnehmen. Setzt geprüfte Paketteile ([Content_Types].xml, *.rels) voraus."""
    vorhanden = set(z.namelist())
    andere = set(_XLSX_FEST)
    mappen = {"xl/workbook.xml"}
    blaetter: set[str] = set()
    try:
        for el in ElementTree.fromstring(z.read("[Content_Types].xml")).iter():
            if _lokal(el.tag) != "Override":
                continue
            teil = el.get("PartName") or ""
            varianten = {teil, teil[1:], teil.lstrip("/")}   # openpyxl: PartName[1:]
            if el.get("ContentType") in _CT_MAPPE:
                mappen |= varianten
            if el.get("ContentType") != _CT_BLATT:
                andere |= varianten
        mappen_rels = {posixpath.join(posixpath.dirname(m), "_rels", posixpath.basename(m) + ".rels")
                       for m in mappen}
        for name in sorted(vorhanden):
            if not name.endswith(".rels"):
                continue
            eltern = posixpath.split(posixpath.dirname(name))[0]
            for el in ElementTree.fromstring(z.read(name)).iter():
                if _lokal(el.tag) != "Relationship" or el.get("TargetMode") == "External":
                    continue
                ziel = el.get("Target") or ""
                ziel = ziel[1:] if ziel.startswith("/") else posixpath.normpath(posixpath.join(eltern, ziel))
                typ = el.get("Type") or ""
                if name in mappen_rels and typ.endswith("/worksheet") and "chartsheet" not in typ:
                    blaetter.add(ziel)
                else:
                    andere.add(ziel)
    except (KeyError, ElementTree.ParseError, ValueError):
        return frozenset()
    return frozenset(b for b in blaetter - andere - mappen if b in vorhanden and not _paketteil(b))


def _xlsx_pruefen(z: zipfile.ZipFile) -> None:
    """Wie `_zip_pruefen` für alle Einträge (welche Teile openpyxl liest, bestimmt die Datei),
    aber streamend gelesene Arbeitsblätter zählen nicht in die Elementsumme: openpyxl liest
    sie im read_only-Modus Zeile für Zeile und hier ohnehin nur bis XLSX_MAX_ZEILEN."""
    infos = z.infolist()
    _zip_groessen(z, infos)
    elemente = _teile_pruefen(z, [i for i in infos if _paketteil(i.filename)])
    _teile_pruefen(z, [i for i in infos if not _paketteil(i.filename)], elemente, _xlsx_blaetter(z))


def _zip_xml(z: zipfile.ZipFile, name: str):
    info = z.getinfo(name)
    if info.file_size > XML_MAX_BYTES:
        raise ValueError(f"{name} ist zu groß")
    return ElementTree.fromstring(z.read(info))


def _w_absatz(p) -> str:
    teile = []
    for el in p.iter():
        if el.tag == _W + "t" and el.text:
            teile.append(el.text)
        elif el.tag == _W + "tab":
            teile.append("\t")
        elif el.tag in (_W + "br", _W + "cr"):
            teile.append("\n")
    return "".join(teile)


_W_STRUKTUR = {_W + "p", _W + "tbl", _W + "tr", _W + "tc"}


def _w_bloecke(element, zeilen: list[str]) -> None:
    """Absätze und Tabellen (Zeile: Zellen mit „ | “ getrennt) in Dokumentreihenfolge."""
    for kind in element:
        if kind.tag == _W + "p":
            zeilen.append(_w_absatz(kind))
        elif kind.tag == _W + "tbl":
            _w_tabelle(kind, zeilen)
        else:  # Inhaltssteuerelemente, Abschnittsgruppen …
            _w_bloecke(kind, zeilen)


def _w_kinder(element, tag: str):
    """Kindelemente `tag`, auch innerhalb von Inhaltssteuerelementen (w:sdt …)."""
    for kind in element:
        if kind.tag == tag:
            yield kind
        elif kind.tag not in _W_STRUKTUR:
            yield from _w_kinder(kind, tag)


def _w_zelle(element, absaetze: list[str], innen: list) -> None:
    for kind in element:
        if kind.tag == _W + "p":
            absaetze.append(_w_absatz(kind).strip())
        elif kind.tag == _W + "tbl":
            innen.append(kind)
        else:
            _w_zelle(kind, absaetze, innen)


def _w_tabelle(tbl, zeilen: list[str]) -> None:
    """Zeilen einer Tabelle; verschachtelte Tabellen folgen als eigene Zeilen. Jedes
    Element wird nur einmal besucht (sonst quadratischer Aufwand bei tiefer Schachtelung)."""
    for tr in _w_kinder(tbl, _W + "tr"):
        zellen, innen = [], []
        for tc in _w_kinder(tr, _W + "tc"):
            absaetze: list[str] = []
            _w_zelle(tc, absaetze, innen)
            zellen.append(" ".join(a for a in absaetze if a))
        if any(zellen):
            zeilen.append(" | ".join(zellen))
        for tabelle in innen:
            _w_tabelle(tabelle, zeilen)


def _teil_nr(name: str) -> int:
    treffer = re.search(r"(\d+)\.xml$", name)
    return int(treffer.group(1)) if treffer else 0


def _docx_auswahl(z: zipfile.ZipFile) -> list[str]:
    """Gelesene Teile: höchstens DOCX_MAX_KOPF_FUSS Kopf- und Fußzeilen, Haupttext."""
    namen = list(dict.fromkeys(z.namelist()))
    kopf = sorted((n for n in namen if re.fullmatch(r"word/header\d*\.xml", n)), key=_teil_nr)
    fuss = sorted((n for n in namen if re.fullmatch(r"word/footer\d*\.xml", n)), key=_teil_nr)
    vorhanden = set(namen)
    return [n for n in (*kopf[:DOCX_MAX_KOPF_FUSS], "word/document.xml",
                        *fuss[:DOCX_MAX_KOPF_FUSS]) if n in vorhanden]


def docx_text(daten: bytes) -> str:
    """Text eines Word-Dokuments: Kopfzeilen (Briefkopf!), Haupttext, Fußzeilen."""
    with zipfile.ZipFile(BytesIO(daten)) as z:
        auswahl = _docx_auswahl(z)
        _zip_pruefen(z, auswahl)
        teile: list[str] = []
        laenge = 0
        for name in auswahl:
            zeilen: list[str] = []
            _w_bloecke(_zip_xml(z, name), zeilen)
            stueck = "\n".join(zeilen).strip()
            if stueck and stueck not in teile:   # gleiche Kopfzeile je Abschnitt nur einmal
                teile.append(stueck)
                laenge += len(stueck)
                if laenge > TEXT_MAX:   # der Rest würde ohnehin abgeschnitten
                    break
    return "\n\n".join(teile)


def _zellwert(wert) -> str:
    if wert is None:
        return ""
    if isinstance(wert, datetime):
        return wert.date().isoformat() if wert.time() == time(0) else wert.isoformat(sep=" ")
    if isinstance(wert, date):
        return wert.isoformat()
    if isinstance(wert, float):
        return str(int(wert)) if wert.is_integer() else f"{wert:.4f}".rstrip("0").rstrip(".")
    if isinstance(wert, Decimal):
        return format(wert, "f")
    return str(wert).strip()


def xlsx_text(daten: bytes) -> tuple[str, int]:
    """Zellwerte je Blatt (Zeilen, Zellen mit „ | “ getrennt) und Anzahl der Blätter."""
    from openpyxl import load_workbook

    with zipfile.ZipFile(BytesIO(daten)) as z:
        _xlsx_pruefen(z)
    mappe = load_workbook(BytesIO(daten), read_only=True, data_only=True, keep_links=False)
    try:
        blaetter = mappe.worksheets
        teile = []
        laenge = 0
        for blatt in blaetter[:XLSX_MAX_BLAETTER]:
            zeilen = []
            for zeile in blatt.iter_rows(max_row=XLSX_MAX_ZEILEN, max_col=XLSX_MAX_SPALTEN,
                                         values_only=True):
                # eine geteilte Zeichenkette kann in jeder Zelle stehen: Länge begrenzen
                werte = [_zellwert(w)[:TEXT_MAX] for w in zeile]
                while werte and not werte[-1]:
                    werte.pop()
                if any(werte):
                    zeilen.append(" | ".join(werte))
                    laenge += len(zeilen[-1]) + 1
                    if laenge > TEXT_MAX:   # der Rest würde ohnehin abgeschnitten
                        break
            if zeilen:
                teile.append(f"Tabellenblatt „{blatt.title}“:\n" + "\n".join(zeilen))
            if laenge > TEXT_MAX:
                break
        return "\n\n".join(teile), len(blaetter)
    finally:
        mappe.close()


# ------------------------------------------------------------------ E-Mail

def _html_text(roh: str) -> str:
    roh = re.sub(r"(?is)<(script|style|head)\b.*?</\1>", " ", roh)
    roh = re.sub(r"(?i)<br\s*/?>|</?(p|div|tr|li|h\d|table)\b[^>]*>", "\n", roh)
    roh = re.sub(r"(?i)</t[dh]>", " | ", roh)
    text = html.unescape(re.sub(r"<[^>]+>", "", roh))
    return "\n".join(re.sub(r"[ \t ]+", " ", z).strip(" |") for z in text.split("\n"))


def _teil_text(teil) -> str:
    try:
        return teil.get_content()
    except Exception:  # noqa: BLE001 – unbekannter Zeichensatz o. Ä.
        nutzlast = teil.get_payload(decode=True) or b""
        return _dekodiere(nutzlast)


def _eml_zerlegen(daten: bytes) -> tuple[list[str], list[str], list[tuple[str, str, bytes]]]:
    """E-Mail → (Kopfzeilen, Textteile, Anhänge [(name, mime, daten)])."""
    nachricht = email.message_from_bytes(daten, policy=email.policy.default)
    kopf = []
    for feld, titel in (("subject", "Betreff"), ("from", "Von"), ("to", "An"), ("date", "Datum")):
        try:
            wert = nachricht.get(feld)
        except Exception:  # noqa: BLE001 – kaputte Kopfzeilen
            wert = None
        if wert:
            kopf.append(f"{titel}: {str(wert).strip()}")
    klartext: list[str] = []
    html_teile: list[str] = []
    anhaenge: list[tuple[str, str, bytes]] = []
    eingebettet: list[tuple[str, str, bytes]] = []   # im HTML-Text gezeigte Bilder (cid:)
    for teil in nachricht.walk():
        if teil.is_multipart():
            continue
        typ = teil.get_content_type()
        name = teil.get_filename()
        art = teil.get_content_disposition()
        if art == "attachment" or (name and not typ.startswith("text/")):
            nutzlast = teil.get_payload(decode=True) or b""
            if not nutzlast:
                continue
            eintrag = (name or "Anhang", typ, nutzlast)
            if (art != "attachment" and teil.get("Content-ID")
                    and ist_bild(mime_bestimmen(nutzlast, typ, name))):
                eingebettet.append(eintrag)
            else:
                anhaenge.append(eintrag)
        elif typ == "text/plain":
            klartext.append(_teil_text(teil))
        elif typ == "text/html":
            html_teile.append(_html_text(_teil_text(teil)))
    # Signatur- und Logobilder nicht als Beleg werten: ohne echte Anhänge zählen eingebettete
    # Bilder ab brauchbarer Größe, daneben nur große (Foto/Screenshot der Bestellung), hinten
    if anhaenge:
        anhaenge += [e for e in eingebettet if _grosses_bild(e[2])]
    else:
        anhaenge = [e for e in eingebettet if not _kleines_bild(e[2])]
    return kopf, (klartext or html_teile), anhaenge


def _office_umfang(daten: bytes, typ: str) -> int:
    """Entpackte Größe der Teile, die das Auslesen liest (DOCX: Auswahl, XLSX: alle)."""
    try:
        with zipfile.ZipFile(BytesIO(daten)) as z:
            if typ == DOCX:
                return sum(z.getinfo(n).file_size for n in _docx_auswahl(z))
            return sum(i.file_size for i in z.infolist())
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        return 0   # nicht lesbar: scheitert gleich beim Öffnen


def _eml(daten: bytes, max_seiten: int, tiefe: int) -> BelegText:
    kopf, texte, anhaenge = _eml_zerlegen(daten)
    teile = ["\n".join(kopf)] if kopf else []
    teile += [t for t in texte if t.strip()]
    warnungen: list[str] = []
    quelle = "TEXT"
    ocr = False
    if len(anhaenge) > EML_MAX_ANHAENGE:
        warnungen.append(f"Nur die ersten {EML_MAX_ANHAENGE} von {len(anhaenge)} Anhängen ausgewertet.")
        anhaenge = anhaenge[:EML_MAX_ANHAENGE]
    # Rechenzeit über alle Anhänge begrenzen: PDF-/Bildseiten (Text bzw. OCR) und entpackte
    # Office-Teile. Doppelt so viele Seiten wie ein Einzeldokument, damit z. B. ein Foto
    # der Bestellung hinter mehrseitigen AGB noch gelesen wird.
    seiten_gesamt = seiten_rest = 2 * max_seiten
    office_rest = ZIP_SUMME_MAX
    for name, typ, inhalt in anhaenge:
        if tiefe >= 1:
            warnungen.append(f"Anhang „{name}“ nicht ausgewertet (verschachtelt).")
            continue
        art = mime_bestimmen(inhalt, typ, name)
        seitenweise = art == PDF or art in BILD_MIME
        if seitenweise and seiten_rest <= 0:
            warnungen.append(f"Anhang „{name}“ nicht ausgewertet: höchstens {seiten_gesamt} Seiten "
                             "aller Anhänge zusammen.")
            continue
        if art in (DOCX, XLSX):
            umfang = _office_umfang(inhalt, art)
            if umfang > office_rest:
                warnungen.append(f"Anhang „{name}“ nicht ausgewertet: Office-Anhänge zusammen entpackt "
                                 f"über {_mb(ZIP_SUMME_MAX)}.")
                continue
            office_rest -= umfang
        unter = _beleg_text(inhalt, art, min(max_seiten, seiten_rest), name, tiefe + 1)
        if seitenweise:
            seiten_rest -= max(1, min(unter.seiten, seiten_rest))
        warnungen += [f"Anhang „{name}“: {w}" for w in unter.warnungen]
        if unter.text.strip():
            teile.append(f"--- Anhang: {name} ---\n{unter.text}")
            ocr = ocr or unter.quelle == "OCR"
    text = _aufraeumen("\n\n".join(teile))
    if ocr:
        quelle = "OCR"
    elif not text.strip():
        quelle = "LEER"
    return BelegText(text, quelle, 1, warnungen)


def eml_anhaenge(daten: bytes) -> list[tuple[str, str, bytes]]:
    """Anhänge einer E-Mail als (Name, MIME-Typ nach Inhalt, Daten)."""
    try:
        _, _, anhaenge = _eml_zerlegen(daten)
    except Exception:  # noqa: BLE001
        return []
    return [(name, mime_bestimmen(inhalt, typ, name), inhalt) for name, typ, inhalt in anhaenge]


# ------------------------------------------------------------------ Einstieg

def _beleg_text(daten: bytes, mime: str, max_seiten: int, dateiname: str, tiefe: int) -> BelegText:
    if not daten:
        return BelegText("", "LEER", 0, ["Die Datei ist leer."])
    typ = mime_bestimmen(daten, mime, dateiname)
    try:
        if typ == PDF:
            erg = _pdf(daten, max_seiten)
        elif typ in BILD_MIME:
            erg = _bild(daten, max_seiten)
        elif typ == DOCX:
            text = _aufraeumen(docx_text(daten))
            erg = BelegText(text, "TEXT" if text.strip() else "LEER", 1, [])
        elif typ == XLSX:
            text, blaetter = xlsx_text(daten)
            text = _aufraeumen(text)
            erg = BelegText(text, "TEXT" if text.strip() else "LEER", blaetter, [])
        elif typ == EML:
            erg = _eml(daten, max_seiten, tiefe)
        elif typ.startswith("text/"):
            roh = _dekodiere(daten)
            text = _aufraeumen(_html_text(roh) if typ == "text/html" else roh)
            erg = BelegText(text, "TEXT" if text.strip() else "LEER", 1, [])
        else:
            return BelegText("", "LEER", 0, [f"Dateityp {typ or 'unbekannt'} kann nicht ausgelesen werden."])
    except Exception as e:  # noqa: BLE001 – nie eine Ausnahme nach außen
        log.warning("Beleg nicht lesbar (%s): %s", typ, e)
        return BelegText("", "LEER", 0, [f"Datei nicht lesbar: {e}"])
    if not erg.text.strip() and erg.quelle != "LEER":
        erg.quelle = "LEER"
    if erg.quelle == "LEER" and not erg.warnungen:
        erg.warnungen.append("Im Dokument wurde kein Text gefunden.")
    if len(erg.text) > TEXT_MAX:
        erg.text = erg.text[:TEXT_MAX]
        erg.warnungen.append(f"Text nach {TEXT_MAX:,} Zeichen abgeschnitten.".replace(",", "."))
    return erg


def beleg_text(daten: bytes, mime: str, max_seiten: int = BELEG_MAX_SEITEN, *,
               dateiname: str = "") -> BelegText:
    """Text eines Belegs. `quelle`: TEXT (Textebene/Datei), OCR (Texterkennung)
    oder LEER (nichts gefunden). Wirft nie – Probleme stehen in `warnungen`."""
    return _beleg_text(daten, mime, max_seiten, dateiname, 0)


def _data_url(bild) -> str:
    return "data:image/png;base64," + base64.b64encode(_png(bild)).decode("ascii")


def vision_bilder(daten: bytes, mime: str, max_seiten: int = VISION_MAX_SEITEN,
                  dpi: int = VISION_DPI, *, dateiname: str = "") -> list[str]:
    """Seitenbilder als data-URLs („data:image/png;base64,…“) für Vision-Modelle:
    PDF-Seiten gerendert, Bilder ausgerichtet und auf max. 2000 px verkleinert,
    bei E-Mails die PDF-/Bildanhänge. Leere Liste bei anderen Typen oder Fehlern."""
    if max_seiten <= 0 or not daten:
        return []
    typ = mime_bestimmen(daten, mime, dateiname)
    try:
        if typ == PDF:
            bilder = pdf_seiten_bilder(daten, max_seiten, dpi)
        elif typ in BILD_MIME:
            bilder = _bild_frames(daten, max_seiten, VISION_MAX_KANTE)
        elif typ == EML:
            # dieselben Anhänge wie beim Text; PDF-Anhänge vor Bildern, Kleinstbilder
            # (Logos, Symbole) auslassen
            anhaenge = [(t, inhalt) for _, t, inhalt in eml_anhaenge(daten)[:EML_MAX_ANHAENGE]
                        if t == PDF or (t in BILD_MIME and not _kleines_bild(inhalt))]
            anhaenge.sort(key=lambda a: a[0] != PDF)
            urls: list[str] = []
            for anhang_typ, inhalt in anhaenge:
                urls += vision_bilder(inhalt, anhang_typ, max_seiten - len(urls), dpi)
                if len(urls) >= max_seiten:
                    break
            return urls[:max_seiten]
        else:
            return []
        return [_data_url(bild_normalisieren(b.convert("RGB") if b.mode not in ("RGB", "L") else b))
                for b in bilder]
    except Exception as e:  # noqa: BLE001
        log.warning("Seitenbilder nicht erzeugbar (%s): %s", typ, e)
        return []
